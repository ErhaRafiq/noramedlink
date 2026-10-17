"""Acceptance tests: route authentication and ownership (Phase 0, extended in Phase 1).

Phase 1 moves routes from Next.js to FastAPI; their checks move with them and call
FastAPI through the same-origin /api/backend proxy, exactly as the frontend does.

Runs against a local Next.js dev server (default http://localhost:3000) and the
local development database. It creates temporary users whose email starts with
"phase0-test-" and deletes them, and only them, together with their related rows,
in a finally block.

Safety:
- Refuses to run unless every database URL in use points to localhost.
  Only the hostname is checked; no URL or secret is printed.
- Never calls external or paid services: no OCR images, no AI requests. Routes
  that would reach OCR.space / OpenAI / Groq are only exercised on paths that
  return before that work (401/403).

Usage (from the repo root, with `npm run dev` running):
    backend\\venv\\Scripts\\python.exe tests\\acceptance_test.py
Optional: PHASE0_BASE_URL=http://localhost:3001 (Next.js), PHASE0_BACKEND_URL=http://localhost:8000 (FastAPI)
"""

from __future__ import annotations

import base64
import hashlib
import hmac
import io
import json
import os
import secrets
import shutil
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx
from PIL import Image
from dotenv import dotenv_values
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from database import DATABASE_URL, get_upload_root_path  # noqa: E402  (backend/database.py; loads .env files first)
from auth import create_access_token  # noqa: E402  (backend/auth.py; reads JWT secret from env)
from config import settings  # noqa: E402  (Stripe webhook secret, to sign test events; never printed)
from services.file_access import _signature  # noqa: E402  (to forge expired links in tests)

BASE_URL = os.environ.get("PHASE0_BASE_URL", "http://localhost:3000").rstrip("/")
# FastAPI itself, for calls that do not go through the Next.js proxy (Stripe webhooks).
BACKEND_URL = os.environ.get("PHASE0_BACKEND_URL", "http://localhost:8000").rstrip("/")
TEST_PREFIX = "phase0-test-"
LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1"}

# Every route patched in Phase 0, with its HTTP method.
# Next.js admin routes were all removed in Phase 1; admin endpoints are in FASTAPI_ADMIN_ROUTES.
ADMIN_ROUTES: list[tuple[str, str]] = []
DOCTOR_ROUTES = [
    ("GET", "/api/backend/doctor/availability"),
    ("PUT", "/api/backend/doctor/availability"),
    ("PUT", "/api/backend/doctor/profile"),
    ("GET", "/api/backend/doctor/dashboard"),
    ("GET", "/api/backend/doctor/overview"),
    ("POST", "/api/backend/doctor/access/unlock"),
    ("POST", "/api/backend/doctor/access/end"),
    ("GET", "/api/backend/doctor/patients/1/clinical-summary"),
    ("POST", "/api/backend/doctor/patients/1/clinical-summary/regenerate"),
    ("GET", "/api/backend/doctor/patients/1/clinical-summary/versions"),
    ("GET", "/api/backend/doctor/patients/1/clinical-summary/versions/1"),
    ("GET", "/api/backend/doctor/medicines/search?q=glucophage"),
    ("POST", "/api/backend/doctor/patients/1/medication-check"),
    ("POST", "/api/backend/doctor/patients/1/prescriptions"),
    ("GET", "/api/backend/doctor/patients/1/prescriptions"),
    ("POST", "/api/backend/doctor/patients/1/prescriptions/1/cancel"),
]
PATIENT_ROUTES = [
]
ANY_USER_ROUTES = [
]
# FastAPI endpoints (through the proxy) that replaced Next routes in Phase 1.
FASTAPI_PATIENT_ROUTES = [
    ("POST", "/api/backend/patient/vitals/basic"),
    ("GET", "/api/backend/patient/vitals/basic"),
    ("GET", "/api/backend/patient/vitals/today"),
    ("GET", "/api/backend/patient/medications"),
    ("POST", "/api/backend/patient/medications"),
    ("POST", "/api/backend/patient/medicine-reminders"),
    ("GET", "/api/backend/patient/medicine-reminders/today"),
    ("PATCH", "/api/backend/patient/medicine-reminders/1/mark-taken"),
    ("PATCH", "/api/backend/patient/medicine-reminders/1/skip"),
    ("POST", "/api/backend/appointments"),
    ("POST", "/api/backend/payments/checkout"),
    ("GET", "/api/backend/payments/session/cs_x"),
    ("POST", "/api/backend/patient/dashboard-otp"),
    ("POST", "/api/backend/emergency/requests"),
    ("GET", "/api/backend/patient/reports"),
    ("GET", "/api/backend/patient/health-flags"),
    ("GET", "/api/backend/patient/prescriptions"),
    ("POST", "/api/backend/appointments/1/review"),
    ("POST", "/api/backend/patient/reports/upload"),
    ("DELETE", "/api/backend/patient/reports/1"),
    ("GET", "/api/backend/patient/reports/1/explanation"),
    ("POST", "/api/backend/patient/reports/1/ask"),
    ("POST", "/api/backend/nlp/summarize/saved-reports"),
]
# FastAPI endpoints restricted to admins.
FASTAPI_ADMIN_ROUTES = [
    ("GET", "/api/backend/api/admin/dashboard-stats"),
    ("GET", "/api/backend/api/admin/users"),
    ("GET", "/api/backend/api/admin/doctors/pending"),
    ("PATCH", "/api/backend/api/admin/doctors/1/approve"),
    ("PATCH", "/api/backend/api/admin/doctors/1/reject"),
    ("GET", "/api/backend/api/admin/recent-activity"),
]
# FastAPI endpoints open to any signed-in role.
FASTAPI_ANY_USER_ROUTES = [
    ("GET", "/api/backend/appointments"),
    ("PUT", "/api/backend/appointments/1"),
    ("GET", "/api/backend/payments"),
    ("POST", "/api/backend/nlp/summarize"),
]
ALL_ROUTES = (ADMIN_ROUTES + DOCTOR_ROUTES + PATIENT_ROUTES + ANY_USER_ROUTES
              + FASTAPI_PATIENT_ROUTES + FASTAPI_ANY_USER_ROUTES + FASTAPI_ADMIN_ROUTES)

# Tables holding rows that belong to a user, children before parents.
# (table, column) pairs; only rows whose column is a phase0-test- user id are deleted.
USER_OWNED_ROWS = [
    ("email_outbox", "user_id"),
    ("auth_tokens", "user_id"),
    ("patient_health_entries", "patient_id"),
    ("audit_logs", "actor_id"),
    ("audit_logs", "patient_id"),
    ("symptom_sessions", "user_id"),
    ("medical_ai_analyses", "user_id"),
    ("notifications", "user_id"),
    ("notification_tokens", "user_id"),
    ("appointment_reviews", "patient_id"),
    ("appointment_reviews", "doctor_id"),
    ("payments", "user_id"),
    ("appointments", "user_id"),
    ("appointments", "doctor_id"),
    ("doctor_availability", "doctor_id"),
    ("emergency_bookings", "user_id"),
    ("emergency_bookings", "doctor_id"),
    ("patient_dashboard_access", "patient_id"),
    ("patient_dashboard_access", "doctor_id"),
    ("patient_dashboard_otps", "patient_id"),
    ("dashboard_otp_attempts", "actor_id"),
    ("reminder_logs", "patient_id"),
    ("medicine_reminders", "patient_id"),
    ("reminders", "user_id"),
    ("vitals", "user_id"),
    ("patient_vitals", "patient_id"),
    ("medications", "user_id"),
    ("medical_documents", "user_id"),
    ("patient_reports", "patient_id"),
    ("medicine_history", "patient_id"),
    ("health_tips", "patient_id"),
    ("health_risk_predictions", "patient_id"),
    ("doctor_summaries", "patient_id"),
    ("clinical_summaries", "patient_id"),
    ("prescriptions", "patient_id"),  # items and overrides go with it (ON DELETE CASCADE)
    ("patient_profiles", "user_id"),
    ("doctor_profiles", "user_id"),
]


# --------------------------------------------------------------------------- safety

def assert_local_databases() -> None:
    """Abort unless every database URL the app uses points to localhost."""
    candidates = {"backend DATABASE_URL in use": DATABASE_URL}
    for env_file in (ROOT / ".env",):  # the project's one env file
        if not env_file.exists():
            continue
        values = dotenv_values(env_file)
        for key in ("DATABASE_URL",):
            if values.get(key):
                candidates[f"{env_file.relative_to(ROOT)} {key}"] = values[key]

    for label, url in candidates.items():
        host = make_url(url.strip().strip('"').strip("'")).host
        if host not in LOCAL_HOSTS:
            sys.exit(f"ABORT: {label} does not point to localhost. Refusing to create test users.")
    print(f"Database guard: {len(candidates)} URL(s) checked, all point to localhost.")
    if settings.resend_api_key:  # blank in a DISABLE_PAID_SERVICES run
        sys.exit(
            "ABORT: an email key is configured. Test accounts use example addresses and tests must never send a "
            "real email. Run the backend and the tests with DISABLE_PAID_SERVICES=1."
        )


# --------------------------------------------------------------------------- results

results: list[tuple[bool, str, str]] = []


def check(name: str, condition: bool, detail: str = "") -> bool:
    results.append((condition, name, detail))
    print(f"  {'PASS' if condition else 'FAIL'}  {name}" + (f"  ({detail})" if detail and not condition else ""))
    return condition


# --------------------------------------------------------------------------- db helpers

engine = create_engine(DATABASE_URL)
run_id = secrets.token_hex(4)


def create_user(label: str, role: str, *, active: bool = True, doctor_status: str = "approved") -> int:
    """Create a test user plus the profile row signup would create for that role."""
    with engine.begin() as conn:
        user_id = conn.execute(
            text(
                """
                insert into users (full_name, email, hashed_password, role, phone,
                                   is_active, is_email_verified, is_identity_verified)
                values (:name, :email, :pw, :role, '', :active, true, true)
                returning id
                """
            ),
            {
                "name": f"{TEST_PREFIX}{label}",
                "email": f"{TEST_PREFIX}{label}-{run_id}@example.test",
                # Not a usable bcrypt hash: these accounts cannot log in with a password.
                "pw": f"{TEST_PREFIX}no-login-{secrets.token_hex(8)}",
                "role": role,
                "active": active,
            },
        ).scalar_one()
        if role == "PATIENT":
            conn.execute(
                text(
                    """
                    insert into patient_profiles (user_id, cnic, date_of_birth, gender, verification_status)
                    values (:uid, '00000-0000000-0', '1990-01-01', 'unspecified', 'verified')
                    """
                ),
                {"uid": user_id},
            )
        elif role == "DOCTOR":
            conn.execute(
                text(
                    """
                    insert into doctor_profiles (user_id, pmdc_number, specialization, hospital_name, verification_status,
                                                 consultation_modes)
                    values (:uid, :pmdc, 'General Physician', 'phase0-test hospital', :status, 'in_clinic,online,home')
                    """
                ),
                {"uid": user_id, "pmdc": f"{TEST_PREFIX}{label}-{run_id}", "status": doctor_status},
            )
            for weekday in range(7):  # open every day 09:00-17:00 in 30-minute slots, so any test can book
                conn.execute(
                    text(
                        "insert into doctor_availability (doctor_id, weekday, start_time, end_time, slot_minutes) "
                        "values (:uid, :d, '09:00', '17:00', 30)"
                    ),
                    {"uid": user_id, "d": weekday},
                )
        return user_id


def create_record(user_id: int, title: str, file_url: str = "") -> int:
    """Seed a report (patient_reports is the single report table since migration 0008)."""
    with engine.begin() as conn:
        return conn.execute(
            text(
                """
                insert into patient_reports (patient_id, title, original_file_name, file_url, file_type,
                                             extracted_text, formatted_text, structured_data_json, summary,
                                             extracted_entities, rule_warnings, validation_status,
                                             validation_results, clinical_summary_json, llm_self_check,
                                             summary_source, ocr_status)
                values (:uid, :title, 'phase0-test.png', :url, 'png', 'Hemoglobin 10.2', 'Hemoglobin 10.2',
                        '{}', 'Advisory summary', '[]', '[]', 'SAFE', '{}', '{}', '{}', 'rule_based', 'COMPLETED')
                returning id
                """
            ),
            {"uid": user_id, "title": title, "url": file_url},
        ).scalar_one()


def create_grant(doctor_id: int, patient_id: int, *, expires_in: timedelta) -> None:
    # Update-then-insert: the live table may lack the (doctor_id, patient_id) unique
    # constraint that the Prisma schema declares, so ON CONFLICT is not usable.
    params = {"d": doctor_id, "p": patient_id, "exp": datetime.now(timezone.utc) + expires_in}
    with engine.begin() as conn:
        updated = conn.execute(
            text(
                """
                update patient_dashboard_access
                set status = 'ACTIVE', expires_at = :exp, ended_at = null, updated_at = now()
                where doctor_id = :d and patient_id = :p
                """
            ),
            params,
        ).rowcount
        if not updated:
            conn.execute(
                text(
                    """
                    insert into patient_dashboard_access (doctor_id, patient_id, status, expires_at)
                    values (:d, :p, 'ACTIVE', :exp)
                    """
                ),
                params,
            )


def configured_paid_keys() -> list[str]:
    """Names of the AI/OCR keys that are set. Tests never call a paid service: they skip instead."""
    if (os.environ.get("DISABLE_PAID_SERVICES") or "").strip().lower() in {"1", "true", "yes"}:
        return []  # the backend ignores every key in this run (config.py), so nothing paid is called
    names = ("AI_API_KEY",)
    return [k for k in names if (os.environ.get(k) or "").strip().strip('"').strip("'")]


def scalar(sql: str, **params):
    with engine.connect() as conn:
        return conn.execute(text(sql), params).scalar()


def cleanup_test_users() -> None:
    """Delete phase0-test- users and rows tied to them. Nothing else is touched."""
    with engine.begin() as conn:
        # "A new doctor is waiting" goes to every real administrator when a test doctor signs up; its text starts
        # with the test doctor's name, so it is removed here (the doctor's own rows go with the doctor).
        conn.execute(text("delete from notifications where kind = 'admin_doctor_review' and body like :p"), {"p": f"{TEST_PREFIX}%"})
        # Symptom checks sent by visitors who were not signed in carry the test marker in their text.
        if conn.execute(text("select to_regclass('public.symptom_sessions')")).scalar():
            conn.execute(text("delete from symptom_sessions where symptom_text like :p"), {"p": f"{TEST_PREFIX}%"})
        ids = list(
            conn.execute(
                text("select id from users where email like :p and full_name like :p"),
                {"p": f"{TEST_PREFIX}%"},
            ).scalars()
        )
        if not ids:
            return
        existing = set(
            conn.execute(text("select table_name from information_schema.tables where table_schema = 'public'")).scalars()
        )
        deleted = 0
        for table, column in USER_OWNED_ROWS:
            if table in existing:
                deleted += conn.execute(text(f"delete from {table} where {column} = any(:ids)"), {"ids": ids}).rowcount
        users = conn.execute(text("delete from users where id = any(:ids)"), {"ids": ids}).rowcount
        # Stripe events signed by the payment checks (ids carry the test prefix; no user column).
        deleted += conn.execute(text("delete from stripe_events where id like :p"), {"p": f"evt_{TEST_PREFIX}%"}).rowcount
    upload_root = get_upload_root_path()
    folders = 0
    for user_id in ids:
        folder = upload_root / str(user_id)
        if folder.is_dir() and folder.parent == upload_root:
            shutil.rmtree(folder)
            folders += 1
    print(f"Cleanup: removed {users} phase0-test- user(s), {deleted} related row(s) and {folders} upload folder(s).")


# --------------------------------------------------------------------------- tokens / http

def token_for(user_id: int, **extra) -> str:
    return create_access_token(str(user_id), extra or None)


def tampered_token(valid_token: str, new_sub: int) -> str:
    """Swap the payload's subject but keep the original signature."""
    header, payload, signature = valid_token.split(".")
    data = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
    data["sub"] = str(new_sub)
    forged = base64.urlsafe_b64encode(json.dumps(data).encode()).decode().rstrip("=")
    return f"{header}.{forged}.{signature}"


client = httpx.Client(base_url=BASE_URL, timeout=180)


def call(method: str, path: str, token: str | None = None, body: dict | None = None) -> httpx.Response:
    headers = {"Authorization": f"Bearer {token}"} if token else {}
    if method == "GET":
        return client.get(path, headers=headers)
    return client.request(method, path, headers=headers, json=body if body is not None else {})


_used_slots: set[str] = set()


def next_slot(doctor_id: int) -> str:
    """The next open slot (UTC, ISO) of a doctor that no earlier check has used, from the public slots endpoint.
    A slot is never handed out twice, so one patient is never booked in two places at once."""
    days = client.get(f"/api/backend/appointments/doctors/{doctor_id}/slots?days=14").json()["days"]
    for day in days:
        for slot in day["slots"]:
            if slot["available"] and slot["start"] not in _used_slots:
                _used_slots.add(slot["start"])
                return slot["start"]
    raise AssertionError(f"doctor {doctor_id} has no open slot left in 14 days")


def overview_patient_ids(doctor_token: str) -> list[int]:
    response = call("GET", "/api/backend/doctor/overview", doctor_token)
    if response.status_code != 200:
        raise AssertionError(f"doctor/overview returned {response.status_code}")
    return [patient["id"] for patient in response.json().get("patients", [])]


# --------------------------------------------------------------------------- tests

def run_tests() -> None:
    patient_a = create_user("patient-a", "PATIENT")
    patient_b = create_user("patient-b", "PATIENT")
    inactive = create_user("patient-inactive", "PATIENT", active=False)
    doctor = create_user("doctor", "DOCTOR")
    doctor_2 = create_user("doctor-2", "DOCTOR")
    admin = create_user("admin", "ADMIN")
    doctor_pending = create_user("doctor-pending", "DOCTOR", doctor_status="pending")
    doctor_inactive = create_user("doctor-inactive", "DOCTOR", active=False)
    doctor_bf = create_user("doctor-bruteforce", "DOCTOR")

    tok_a, tok_b = token_for(patient_a), token_for(patient_b)
    tok_doctor, tok_doctor_2, tok_admin = token_for(doctor), token_for(doctor_2), token_for(admin)

    # Legacy identity fields that must now be ignored.
    spoof_body = {"userId": admin, "adminId": admin, "doctorId": doctor, "patientId": patient_a}

    print("\n[0] /api/backend proxy only reaches the configured FastAPI server")
    check("proxy forwards to FastAPI", call("GET", "/api/backend/health").status_code == 200)
    response = call("GET", "/api/backend/health%00")
    check("control character in path -> 400", response.status_code == 400, f"got {response.status_code}")
    for path in ("/api/backend/http:%2F%2Fexample.com/x", "/api/backend/%2F%2Fexample.com/x", "/api/backend/@example.com/x"):
        response = call("GET", path)
        detail = response.json().get("detail") if "json" in response.headers.get("content-type", "") else None
        check(f"{path} stays on FastAPI (its own 404)", response.status_code == 404 and detail == "Not Found",
              f"got {response.status_code} {detail!r}")

    print("\n[1] No token -> 401 on every patched route (legacy userId params ignored)")
    for method, path in ALL_ROUTES:
        target = f"{path}?userId={admin}" if method == "GET" else path
        response = call(method, target, None, spoof_body)
        check(f"{method} {path} without token -> 401", response.status_code == 401, f"got {response.status_code}")

    print("\n[2] Invalid tokens -> 401")
    response = call("GET", "/api/backend/patient/reports", token_for(patient_a, exp=int(time.time()) - 60))
    check("expired token -> 401", response.status_code == 401, f"got {response.status_code}")
    response = call("GET", "/api/backend/api/admin/dashboard-stats", tampered_token(tok_a, admin))
    check("token with tampered subject -> 401", response.status_code == 401, f"got {response.status_code}")
    response = call("GET", "/api/backend/patient/reports", "not-a-jwt")
    check("malformed token -> 401", response.status_code == 401, f"got {response.status_code}")

    print("\n[3] Wrong role -> 403")
    for method, path in ADMIN_ROUTES + FASTAPI_ADMIN_ROUTES:
        response = call(method, path, tok_a)
        check(f"patient on {path} -> 403", response.status_code == 403, f"got {response.status_code}")
        response = call(method, path, tok_doctor)
        check(f"doctor on {path} -> 403", response.status_code == 403, f"got {response.status_code}")
    for method, path in DOCTOR_ROUTES:
        response = call(method, path, tok_a)
        check(f"patient on {path} -> 403", response.status_code == 403, f"got {response.status_code}")
    for method, path in PATIENT_ROUTES + FASTAPI_PATIENT_ROUTES:
        response = call(method, path, tok_doctor)
        check(f"doctor on {path} -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("POST", "/api/backend/nlp/summarize/saved-reports", tok_doctor)
    check("doctor on summarize saved-report mode -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("GET", "/api/backend/patient/reports", token_for(inactive))
    check("inactive patient account -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("GET", "/api/backend/patient/vitals/today", token_for(inactive))
    check("inactive patient account on vitals -> 403", response.status_code == 403, f"got {response.status_code}")

    print("\n[4] Correct role -> allowed (positive controls)")
    check("admin on admin dashboard-stats -> 200", call("GET", "/api/backend/api/admin/dashboard-stats", tok_admin).status_code == 200)
    check("doctor on doctor/overview -> 200", call("GET", "/api/backend/doctor/overview", tok_doctor).status_code == 200)
    check("patient on reports -> 200", call("GET", "/api/backend/patient/reports", tok_a).status_code == 200)

    print("\n[5] Ownership: Patient A cannot read or delete Patient B's records")
    record_a = create_record(patient_a, f"{TEST_PREFIX}record A")
    record_b = create_record(patient_b, f"{TEST_PREFIX}record B")

    response = call("GET", f"/api/backend/patient/reports?user_id={patient_b}&userId={patient_b}", tok_a)
    record_ids = [record["id"] for record in response.json().get("items", [])]
    check("A's report list has A's record", record_a in record_ids, f"ids={record_ids}")
    check("A's report list (with ?userId=B) does not list B's record", record_b not in record_ids, f"ids={record_ids}")

    response = call("DELETE", f"/api/backend/patient/reports/{record_b}?user_id={patient_b}", tok_a)
    check("A deleting B's report -> 404", response.status_code == 404, f"got {response.status_code}")
    check("B's report still exists after A's attempt",
          scalar("select count(*) from patient_reports where id = :id", id=record_b) == 1)
    response = call("GET", f"/api/backend/patient/reports/{record_b}", tok_a)
    check("A reading B's report -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("GET", f"/api/backend/patient/reports/{record_b}/explanation", tok_a)
    check("A reading B's report explanation -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("POST", f"/api/backend/patient/reports/{record_b}/ask", tok_a, {"question": "What does this mean?"})
    check("A asking about B's report -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("GET", f"/api/backend/patient/reports/{record_a}/explanation", tok_a)
    explained = response.json() if response.status_code == 200 else {}
    check("A's own report explanation -> 200 with the machine-generated note",
          response.status_code == 200 and "not by a doctor" in explained.get("disclaimer", ""), f"got {response.status_code}")
    response = call("POST", f"/api/backend/patient/reports/{record_a}/ask", tok_a, {"question": "I have chest pain"})
    check("A asking an emergency question -> the 1122 message",
          response.status_code == 200 and response.json().get("emergency") is True and "1122" in response.json().get("answer", ""),
          f"got {response.status_code}")
    check("A question that is too short -> 422",
          call("POST", f"/api/backend/patient/reports/{record_a}/ask", tok_a, {"question": "hi"}).status_code == 422)

    response = call("DELETE", f"/api/backend/patient/reports/{record_b}", tok_b)
    check("B deleting own report -> 200", response.status_code == 200, f"got {response.status_code}")
    check("B's report is gone", scalar("select count(*) from patient_reports where id = :id", id=record_b) == 0)

    response = call("POST", "/api/backend/patient/vitals/basic", tok_a,
                    {"user_id": patient_b, "userId": patient_b, "blood_pressure": "120/80", "sugar_level": "95"})
    check("A adds vitals (body user id=B ignored) -> 201", response.status_code == 201, f"got {response.status_code}")
    check("vitals stored under A", scalar("select count(*) from vitals where user_id = :id", id=patient_a) == 1)
    check("no vitals stored under B", scalar("select count(*) from vitals where user_id = :id", id=patient_b) == 0)
    response = call("GET", f"/api/backend/patient/vitals/basic?user_id={patient_b}&userId={patient_b}", tok_a)
    history = response.json() if response.status_code == 200 else []
    check("vitals/basic history (with ?user_id=B) lists only A's entry",
          len(history) == 1 and history[0].get("user_id") == patient_a, f"got {response.status_code} {history}")
    response = call("GET", "/api/backend/patient/vitals/basic", tok_b)
    check("B's vitals/basic history is empty", response.status_code == 200 and response.json() == [])
    response = call("GET", "/api/backend/patient/vitals/today", tok_b)
    check("B's vitals/today is empty", response.status_code == 200 and response.json() is None)

    print("\n[5b] Vitals: nothing is estimated or invented")
    today = call("GET", "/api/backend/patient/vitals/today", tok_a).json() or {}
    check("simple log only -> BP parsed from it", (today.get("blood_pressure_systolic"), today.get("blood_pressure_diastolic")) == (120, 80), f"got {today}")
    check("simple log only -> heart rate, SpO2, temperature are null (not 72/98/98.4)",
          today.get("heart_rate") is None and today.get("spo2") is None and today.get("temperature") is None, f"got {today}")

    print("\n[5c] Medication schedules are per patient")
    med = {"medicineName": "Phase0 Testamol", "dosage": "1 tablet", "startDate": "2026-01-01", "userId": patient_b}
    response = call("POST", "/api/backend/patient/medications", tok_a, med)
    created = response.json() if response.status_code == 201 else {}
    check("A creates a medication (body userId=B ignored) -> 201", response.status_code == 201, f"got {response.status_code}")
    check("medication stored under A", created.get("patient_id") == patient_a
          and scalar("select user_id from medications where id = :id", id=created.get("id", 0)) == patient_a)
    check("A's list has it", [m["id"] for m in call("GET", "/api/backend/patient/medications", tok_a).json()] == [created.get("id")])
    check("B's list is empty", call("GET", "/api/backend/patient/medications", tok_b).json() == [])
    response = call("POST", "/api/backend/patient/medications", tok_a, {**med, "endDate": "2025-01-01"})
    check("end date before start date -> 400", response.status_code == 400, f"got {response.status_code}")

    print("\n[5d] Reminders are per patient")
    due = datetime.now(timezone.utc).isoformat()
    med_id = created.get("id")
    schedule = {"medicationId": med_id, "reminderTime": due, "frequency": "Once daily"}
    response = call("POST", "/api/backend/patient/medicine-reminders", tok_b, schedule)
    check("B scheduling a dose of A's medication -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("POST", "/api/backend/patient/medicine-reminders", tok_a, schedule)
    dose = response.json() if response.status_code == 201 else {}
    check("A schedules a dose -> 201, due now", response.status_code == 201 and dose.get("status") == "DUE_NOW",
          f"got {response.status_code} {dose.get('status')}")
    check("dose includes its medication", (dose.get("medication") or {}).get("medicine_name") == "Phase0 Testamol")
    check("scheduling notification has no reminder_id (it belongs to the other reminders table)",
          scalar("select count(*) from notifications where user_id = :id and title = 'Medicine reminder scheduled' "
                 "and reminder_id is null", id=patient_a) == 1)
    check("A's today list has the dose", [r["id"] for r in call("GET", "/api/backend/patient/medicine-reminders/today", tok_a).json()] == [dose.get("id")])
    check("B's today list is empty", call("GET", "/api/backend/patient/medicine-reminders/today", tok_b).json() == [])
    response = call("PATCH", f"/api/backend/patient/medicine-reminders/{dose.get('id')}/mark-taken", tok_b)
    check("B marking A's dose taken -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("PATCH", f"/api/backend/patient/medicine-reminders/{dose.get('id')}/mark-taken", tok_a)
    taken = response.json() if response.status_code == 200 else {}
    check("A marks the dose taken -> TAKEN with a time", taken.get("status") == "TAKEN" and bool(taken.get("taken_at")), f"got {taken}")
    check("reminder log has the scheduled and taken entries",
          scalar("select count(*) from reminder_logs where reminder_id = :id", id=dose.get("id", 0)) == 2)
    response = call("PATCH", f"/api/backend/patient/medicine-reminders/{dose.get('id')}/skip", tok_a)
    skipped = response.json() if response.status_code == 200 else {}
    check("A skips it -> SKIPPED with a time", skipped.get("status") == "SKIPPED" and bool(skipped.get("skipped_at")),
          f"got {response.status_code} {skipped.get('status')}")

    print("\n[5e] Doctor search and booking: verified doctors only, nothing invented")
    listed = call("GET", "/api/backend/appointments/doctors").json()
    listed_ids = {d["doctor_id"] for d in listed}
    check("approved doctor is listed", doctor in listed_ids)
    check("pending (unverified) doctor is not listed", doctor_pending not in listed_ids)
    check("inactive doctor is not listed", doctor_inactive not in listed_ids)
    check("no demo doctors (every entry is a registered doctor)", all(d["source"] == "registered" and d["doctor_id"] for d in listed))
    check("no invented ratings (nobody has been reviewed yet)", all(d["rating"] is None and d["rating_count"] == 0 for d in listed))

    booking = {"doctor_id": doctor, "doctor_name": "Someone Else", "specialty": "Cardiology", "date": next_slot(doctor),
               "reason": "Follow-up", "user_id": patient_b}
    response = call("POST", "/api/backend/appointments", tok_a, booking)
    booked = (response.json() or {}).get("appointment", {}) if response.status_code == 201 else {}
    check("A books the approved doctor -> 201", response.status_code == 201, f"got {response.status_code}")
    check("booking stored under A with the doctor's real name and specialty (request values ignored)",
          booked.get("user_id") == patient_a and booked.get("doctor_name") == f"{TEST_PREFIX}doctor"
          and booked.get("specialty") == "General Physician" and booked.get("reason") == "Follow-up", f"got {booked}")
    response = call("POST", "/api/backend/appointments", tok_a, {**booking, "doctor_id": doctor_pending})
    check("booking an unverified doctor -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("POST", "/api/backend/appointments", tok_a, {"doctor_name": f"{TEST_PREFIX}doctor", "date": next_slot(doctor)})
    check("booking by name only (no doctor_id) -> 422", response.status_code == 422, f"got {response.status_code}")

    check("A sees the appointment", [a["id"] for a in call("GET", "/api/backend/appointments", tok_a).json()["items"]] == [booked.get("id")])
    check("B sees no appointments", call("GET", "/api/backend/appointments", tok_b).json()["items"] == [])
    check("the booked doctor sees it", booked.get("id") in [a["id"] for a in call("GET", "/api/backend/appointments", tok_doctor).json()["items"]])
    check("another doctor does not", call("GET", "/api/backend/appointments", tok_doctor_2).json()["items"] == [])

    appt = f"/api/backend/appointments/{booked.get('id')}"
    response = call("PUT", appt, tok_a, {"payment_status": "SUCCESS"})
    check("patient marking own appointment paid -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("PUT", appt, tok_a, {"status": "COMPLETED"})
    check("patient setting status COMPLETED -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("PUT", appt, tok_b, {"status": "CANCELLED"})
    check("B changing A's appointment -> 403", response.status_code == 403, f"got {response.status_code}")
    later = datetime.fromisoformat(next_slot(doctor).replace("Z", "+00:00"))
    response = call("PUT", appt, tok_a, {"date": later.isoformat()})
    check("patient reschedules own appointment (date only) -> 200",
          response.status_code == 200 and response.json().get("status") == "BOOKED", f"got {response.status_code}")
    clinic_local = later.astimezone(timezone(timedelta(hours=5))).replace(tzinfo=None)  # clinic time, no zone
    response = call("PUT", appt, tok_a, {"date": clinic_local.isoformat()})
    check("a date without a time zone is read as clinic time (was a 500)", response.status_code == 200, f"got {response.status_code}")
    response = call("PUT", appt, tok_a, {"date": later.isoformat(), "status": "BOOKED"})
    check("patient setting status BOOKED is still refused -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("PUT", appt, tok_a, {"status": "CANCELLED"})
    check("patient can cancel own appointment -> 200", response.status_code == 200 and response.json().get("status") == "CANCELLED",
          f"got {response.status_code}")

    print("\n[5f] Payments: the doctor sets the fee, Stripe confirms; nobody marks a payment paid by hand")
    doctor_form = {"hospital_name": "Test Hospital", "city": "Lahore", "qualifications": "MBBS, FCPS", "experience_years": 8,
                   "bio": "Phase0 test doctor", "languages": ["Urdu", "English"], "consultation_modes": ["in_clinic", "online"],
                   "consultation_fee": 1500}
    response = call("PUT", "/api/backend/doctor/profile", tok_doctor, doctor_form)
    check("doctor saves a profile with a fee -> 200", response.status_code == 200
          and response.json()["doctor"]["consultation_fee"] == 1500, f"got {response.status_code}")
    response = call("PUT", "/api/backend/doctor/profile", tok_doctor, {**doctor_form, "consultation_fee": 50})
    check("a fee under PKR 200 -> 422", response.status_code == 422, f"got {response.status_code}")
    response = call("PUT", "/api/backend/doctor/profile", tok_a, doctor_form)
    check("a patient saving a doctor profile -> 403", response.status_code == 403, f"got {response.status_code}")
    listed = {d["doctor_id"]: d for d in call("GET", "/api/backend/appointments/doctors").json()}
    check("the directory shows the doctor's fee", (listed.get(doctor) or {}).get("consultation_fee") == 1500)

    response = call("POST", "/api/backend/appointments", tok_a, {"doctor_id": doctor, "date": next_slot(doctor), "reason": "Payment test"})
    to_pay = (response.json() or {}).get("appointment", {}) if response.status_code == 201 else {}
    check("a new appointment carries the fee and is unpaid",
          to_pay.get("consultation_fee") == 1500 and to_pay.get("payment_status") == "PENDING", f"got {to_pay}")
    response = call("POST", "/api/backend/payments/checkout", tok_b, {"appointment_id": to_pay.get("id")})
    check("B opening checkout for A's appointment -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("POST", "/api/backend/payments/checkout", tok_doctor, {"appointment_id": to_pay.get("id")})
    check("a doctor opening checkout -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("POST", "/api/backend/payments/checkout", tok_a, {"appointment_id": booked.get("id")})
    check("checkout for a cancelled appointment -> 409", response.status_code == 409, f"got {response.status_code}")
    if not settings.stripe_secret_key:
        response = call("POST", "/api/backend/payments/checkout", tok_a, {"appointment_id": to_pay.get("id")})
        check("checkout without a Stripe key -> 503 and no payment recorded", response.status_code == 503
              and scalar("select count(*) from payments where appointment_id = :id", id=to_pay.get("id", 0)) == 0,
              f"got {response.status_code}")
    response = call("PUT", f"/api/backend/appointments/{to_pay.get('id')}", tok_admin, {"payment_status": "PAID"})
    check("even an admin cannot mark an appointment paid -> 403", response.status_code == 403, f"got {response.status_code}")

    if not settings.stripe_webhook_secret:
        response = call("POST", "/api/backend/payments/stripe/webhook", None, {})
        check("webhook without STRIPE_WEBHOOK_SECRET set -> 503", response.status_code == 503, f"got {response.status_code}")
        print("  SKIP  signed webhook checks (set STRIPE_WEBHOOK_SECRET to run them; no Stripe call is made)")
    else:
        # Stripe calls FastAPI directly (the Next.js proxy does not pass Stripe-Signature on).
        backend = httpx.Client(base_url=BACKEND_URL, timeout=30)
        session_id = f"cs_{TEST_PREFIX}{run_id}"
        with engine.begin() as conn:
            conn.execute(text(
                "insert into payments (user_id, appointment_id, amount, currency, method, purpose, status, reference, "
                "stripe_session_id) values (:u, :a, 1500, 'pkr', 'STRIPE_CARD', 'Consultation', 'PENDING', :r, :s)"),
                {"u": patient_a, "a": to_pay.get("id"), "r": f"{TEST_PREFIX}{run_id}", "s": session_id})

        def stripe_event(event_id: str, kind: str, **session) -> tuple[bytes, dict]:
            body = json.dumps({"id": f"evt_{TEST_PREFIX}{run_id}-{event_id}", "type": kind, "data": {"object": {
                "id": session_id, "payment_status": "paid", "amount_total": 150_000, "currency": "pkr",
                "payment_intent": "pi_test", **session}}}).encode()
            stamp = int(time.time())
            digest = hmac.new(settings.stripe_webhook_secret.encode(), f"{stamp}.".encode() + body, hashlib.sha256).hexdigest()
            return body, {"Stripe-Signature": f"t={stamp},v1={digest}", "Content-Type": "application/json"}

        def paid_status() -> str:
            return scalar("select status from payments where stripe_session_id = :s", s=session_id)

        body, headers = stripe_event("forged", "checkout.session.completed")
        response = backend.post("/payments/stripe/webhook", content=body, headers={**headers, "Stripe-Signature": "t=1,v1=00"})
        check("a webhook with a forged signature -> 400, nothing paid", response.status_code == 400 and paid_status() == "PENDING",
              f"got {response.status_code}")
        body, headers = stripe_event("short", "checkout.session.completed", amount_total=100)
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("a paid session for the wrong amount does not mark it paid",
              response.status_code == 200 and response.json().get("outcome") == "amount mismatch" and paid_status() == "PENDING",
              f"got {response.status_code}")
        body, headers = stripe_event("paid", "checkout.session.completed")
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("a signed checkout.session.completed marks the payment PAID",
              response.status_code == 200 and response.json().get("outcome") == "paid" and paid_status() == "PAID",
              f"got {response.status_code}")
        check("and the appointment PAID",
              scalar("select payment_status from appointments where id = :id", id=to_pay.get("id", 0)) == "PAID")
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("the same event delivered again changes nothing (duplicate)", response.json().get("outcome") == "duplicate")
        check("the paid payment is told to the patient once, and to the doctor once, however often the webhook arrives",
              scalar("select count(*) from notifications where user_id = :u and kind = 'payment_received' and link = '/payments'", u=patient_a) == 1
              and scalar("select count(*) from notifications where user_id = :d and kind = 'payment_received'", d=doctor) == 1)
        body, headers = stripe_event("late-expiry", "checkout.session.expired")
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("a late 'expired' event never undoes a paid payment",
              response.json().get("outcome") == "already paid" and paid_status() == "PAID")
        response = call("GET", f"/api/backend/payments/session/{session_id}", tok_a)
        check("A reads the paid payment for the result page", response.status_code == 200 and response.json().get("status") == "PAID")
        response = call("GET", f"/api/backend/payments/session/{session_id}", tok_b)
        check("B reading A's payment -> 404", response.status_code == 404, f"got {response.status_code}")
        response = call("POST", "/api/backend/payments/checkout", tok_a, {"appointment_id": to_pay.get("id")})
        check("checkout for a paid appointment -> 409", response.status_code == 409, f"got {response.status_code}")

        payment_id = scalar("select id from payments where stripe_session_id = :s", s=session_id)
        check("a receipt without a token -> 401", call("GET", f"/api/backend/payments/{payment_id}").status_code == 401)
        response = call("GET", f"/api/backend/payments/{payment_id}", tok_a)
        check("A reads their own receipt", response.status_code == 200 and response.json().get("id") == payment_id,
              f"got {response.status_code}")
        check("B reading A's receipt -> 404", call("GET", f"/api/backend/payments/{payment_id}", tok_b).status_code == 404)
        response = call("PUT", f"/api/backend/appointments/{to_pay.get('id')}", tok_a, {"status": "CANCELLED"})
        check("A cancels the paid appointment -> 200", response.status_code == 200, f"got {response.status_code}")
        # No Stripe key in tests: the refund cannot be requested, and that must be visible, never 'refunded'.
        check("without Stripe the refund shows as needing attention, not refunded",
              paid_status() == "REFUND_FAILED"
              and scalar("select payment_status from appointments where id = :id", id=to_pay.get("id", 0)) == "REFUND_FAILED")
        check("A is told about the refund", scalar(
            "select count(*) from notifications where user_id = :u and kind = 'payment_refund'", u=patient_a) == 1)
        body, headers = stripe_event("part-refund", "charge.refunded", amount_refunded=100)
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("a partial refund event does not mark it refunded",
              response.json().get("outcome") == "partial refund ignored" and paid_status() == "REFUND_FAILED")
        body, headers = stripe_event("refund", "charge.refunded", amount_refunded=150_000)
        response = backend.post("/payments/stripe/webhook", content=body, headers=headers)
        check("a signed charge.refunded for the full amount marks it REFUNDED",
              response.json().get("outcome") == "refunded" and paid_status() == "REFUNDED"
              and scalar("select refunded_at is not null from payments where stripe_session_id = :s", s=session_id) is True)
        check("a completed refund is told to the patient once", scalar("select count(*) from notifications where user_id = :u and kind = 'payment_refunded'", u=patient_a) == 1)
        check("and the appointment REFUNDED",
              scalar("select payment_status from appointments where id = :id", id=to_pay.get("id", 0)) == "REFUNDED")
        backend.close()
    check("B sees no payments", call("GET", "/api/backend/payments", tok_b).json()["items"] == [])

    print("\n[5g] Home-visit requests: no invented dispatch")
    request = {"patientName": "Phase0 A", "phone": "03000000000", "address": "Test street", "symptoms": "fever",
               "userId": patient_b, "user_id": patient_b}
    response = call("POST", "/api/backend/emergency/requests", tok_a, request)
    body = response.json() if response.status_code == 201 else {}
    booking_out = body.get("booking") or {}
    check("A sends a request -> 201 REQUESTED", response.status_code == 201 and booking_out.get("status") == "REQUESTED",
          f"got {response.status_code} {booking_out.get('status')}")
    check("no doctor and no ETA are claimed", booking_out.get("doctorName") is None and booking_out.get("estimatedMinutes") is None,
          f"got {booking_out}")
    check("response directs emergencies to 1122", "1122" in body.get("message", "") and "1122" in body.get("emergencyNotice", ""))
    check("request stored under A (body user id ignored), no doctor",
          scalar("select count(*) from emergency_bookings where user_id = :a and doctor_id is null and status = 'REQUESTED'",
                 a=patient_a) == 1
          and scalar("select count(*) from emergency_bookings where user_id = :b", b=patient_b) == 0)

    response = call("POST", "/api/backend/nlp/summarize", tok_a, {"text": "too short"})
    check("summarize text under 20 characters -> 400", response.status_code == 400, f"got {response.status_code}")
    response = call("POST", "/api/backend/nlp/summarize", tok_a, {"text": "x" * 20001})
    check("summarize text over 20,000 characters -> 413", response.status_code == 413, f"got {response.status_code}")
    response = call("POST", "/api/ai/chat", tok_a, {"message": "hello"})
    check("AI chat route removed (rebuilt in Phase 2a) -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("GET", "/api/backend/patient/documents", tok_a)
    check("removed /patient/documents endpoint -> 404", response.status_code == 404, f"got {response.status_code}")

    print("\n[5i] Symptom checker: open to visitors, red flags block everything else")
    probe = f"{TEST_PREFIX}{run_id} "
    check_path = "/api/backend/symptoms/check"
    response = call("POST", check_path, None, {"text": probe + "Chest pain, dizziness, shortness of breath"})
    emergency = response.json() if response.status_code == 200 else {}
    check("a visitor with chest pain and breathlessness gets the emergency notice (no login needed)",
          response.status_code == 200 and emergency.get("emergency") is True and "1122" in (emergency.get("emergency_notice") or ""),
          f"got {response.status_code}")
    check("an emergency offers no specialist and no doctor to book",
          emergency.get("suggestions") == [] and len(emergency.get("warning_signs", [])) >= 2)
    response = call("POST", check_path, None, {"text": probe + "itchy rash on arm for a week"})
    rash = response.json() if response.status_code == 200 else {}
    check("an itchy rash -> not an emergency, Dermatologist first",
          response.status_code == 200 and rash.get("emergency") is False and rash.get("emergency_notice") is None
          and rash["suggestions"][0]["specialty"] == "Dermatologist", f"got {response.status_code}")
    check("the advice is marked advisory", "not a diagnosis" in rash.get("advisory", ""))
    response = call("POST", check_path, tok_a, {"text": probe + "fever and sore throat for 2 days"})
    check("fever and a sore throat -> General Physician", response.status_code == 200
          and response.json()["suggestions"][0]["specialty"] == "General Physician", f"got {response.status_code}")
    check("the suggested doctors never include an unverified doctor", all(
        d["id"] != doctor_pending for s_ in response.json()["suggestions"] for d in s_["doctors"]))
    check("a signed-in check is saved under that patient",
          scalar("select count(*) from symptom_sessions where user_id = :a", a=patient_a) == 1)
    check("a visitor's checks are saved without a user",
          scalar("select count(*) from symptom_sessions where symptom_text like :p and user_id is null", p=probe + "%") == 2)
    check("a red-flag run is stored as an emergency with no specialty",
          scalar("select red_flag and specialty is null from symptom_sessions where symptom_text like :p", p=probe + "Chest%") is True)
    check("a one-word text -> 422", call("POST", check_path, None, {"text": "a"}).status_code == 422)
    check("a text over 1,000 characters -> 422", call("POST", check_path, None, {"text": "x" * 1001}).status_code == 422)
    check("a user id in the body is ignored (identity comes from the token)",
          call("POST", check_path, tok_a, {"text": probe + "back pain", "user_id": patient_b}).status_code == 200
          and scalar("select count(*) from symptom_sessions where user_id = :b", b=patient_b) == 0)
    check("admin symptom stats without a token -> 401",
          call("GET", "/api/backend/api/admin/symptom-stats").status_code == 401)
    check("a patient on admin symptom stats -> 403",
          call("GET", "/api/backend/api/admin/symptom-stats", tok_a).status_code == 403)
    stats = call("GET", "/api/backend/api/admin/symptom-stats", tok_admin)
    stats_body = stats.json() if stats.status_code == 200 else {}
    check("admin reads counts, never the text", stats.status_code == 200 and stats_body.get("emergencies", 0) >= 1
          and stats_body.get("total", 0) >= 4 and probe not in stats.text, f"got {stats.status_code}")

    print("\n[5h] Report upload (browser-extracted text, no OCR.space)")
    paid_keys = configured_paid_keys()
    if paid_keys:
        print("  SKIP  upload pipeline test: an AI/OCR key is configured, and tests never call paid services")
    else:
        png = io.BytesIO()
        Image.new("RGB", (8, 8), "white").save(png, format="PNG")
        response = client.post(
            "/api/backend/patient/reports/upload",
            headers={"Authorization": f"Bearer {tok_a}"},
            data={"title": "Phase0 CBC", "extracted_text": "Hemoglobin: 10.2 g/dL\nGlucose: 250 mg/dL"},
            files={"file": ("phase0-cbc.png", png.getvalue(), "image/png")},
        )
        uploaded = response.json() if response.status_code == 201 else {}
        check("A uploads a report with browser-extracted text -> 201 COMPLETED",
              response.status_code == 201 and uploaded.get("ocr_status") == "COMPLETED", f"got {response.status_code}")
        check("file stored under A's upload folder, not as base64",
              str(uploaded.get("file_url", "")).startswith(f"uploads/{patient_a}/"), f"got {uploaded.get('file_url')}")
        import docx as word

        document = word.Document()
        document.add_paragraph("Hemoglobin 10.2 g/dL")
        table = document.add_table(rows=1, cols=3)
        for cell_, value in zip(table.rows[0].cells, ("Glucose", "250", "mg/dL"), strict=True):
            cell_.text = value
        word_file = io.BytesIO()
        document.save(word_file)
        response = client.post(
            "/api/backend/patient/reports/upload",
            headers={"Authorization": f"Bearer {tok_a}"},
            data={"title": "Phase0 Word report"},
            files={"file": ("phase0.docx", word_file.getvalue(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        word_report = response.json() if response.status_code == 201 else {}
        check("A uploads a Word report (read on the server, no browser text) -> 201 with its text and table",
              response.status_code == 201 and "Glucose 250 mg/dL" in (word_report.get("extracted_text") or ""),
              f"got {response.status_code} {str(word_report.get('extracted_text'))[:80]!r}")
        response = client.post(
            "/api/backend/patient/reports/upload",
            headers={"Authorization": f"Bearer {tok_a}"},
            data={"title": "Phase0 legacy Word"},
            files={"file": ("old.doc", b"x", "application/msword")},
        )
        check("a legacy .doc file is refused -> 400", response.status_code == 400, f"got {response.status_code}")
        listed = call("GET", "/api/backend/patient/reports", tok_a).json()["items"]
        check("upload appears in A's report list", uploaded.get("id") in [r["id"] for r in listed])
        response = call("DELETE", f"/api/backend/patient/reports/{uploaded.get('id')}", tok_b)
        check("B deleting A's upload -> 404", response.status_code == 404, f"got {response.status_code}")
        # B's report gets a value A's reports do not have; A's summary must never contain it.
        with engine.begin() as conn:
            conn.execute(text("update patient_reports set extracted_text = 'Creatinine 9.9 mg/dL', "
                              "formatted_text = 'Creatinine 9.9 mg/dL' where id = :id"), {"id": record_b})
        response = call("POST", "/api/backend/nlp/summarize/saved-reports", tok_a)
        saved = response.json() if response.status_code == 200 else {}
        findings = saved.get("key_findings", [])
        check("saved-report summary uses only A's reports",
              response.status_code == 200 and "Hemoglobin: 10.2" in findings
              and not any("9.9" in item for item in findings + saved.get("risk_indicators", []))
              and bool(saved.get("summary")), f"got {response.status_code} {findings}")
        check("saved-report summary no longer returns the combined report text", "text" not in saved)
        response = call("POST", "/api/backend/nlp/summarize", tok_doctor, {"text": "Hemoglobin 10.2 g/dL, glucose 250 mg/dL."})
        check("doctor can summarize pasted text -> 200", response.status_code == 200 and bool(response.json().get("summary")),
              f"got {response.status_code}")

    print("\n[6] Doctor grant lifecycle on doctor/overview")
    profile = call("GET", "/api/backend/doctor/dashboard", tok_doctor).json()
    check("doctor/dashboard returns only the doctor's own profile (no placeholder stats)",
          set(profile) == {"doctor"} and profile["doctor"]["id"] == doctor, f"got {profile}")
    check("no grant -> patient A not listed", patient_a not in overview_patient_ids(tok_doctor))

    create_grant(doctor, patient_a, expires_in=timedelta(minutes=-5))
    check("expired grant -> patient A not listed", patient_a not in overview_patient_ids(tok_doctor))

    overview = call("GET", "/api/backend/doctor/overview", tok_doctor).json()
    check("no grant -> none of A's medical records or medications visible",
          all(r["userId"] != patient_a for r in overview["medicalRecords"] + overview["medications"]))
    create_grant(doctor, patient_a, expires_in=timedelta(hours=1))
    check("active grant -> patient A listed", patient_a in overview_patient_ids(tok_doctor))
    overview = call("GET", "/api/backend/doctor/overview", tok_doctor).json()
    granted_record = next((r for r in overview["medicalRecords"] if r["id"] == record_a), {})
    check("doctor view reads the report table (department, summary, status)",
          granted_record.get("summary") == "Advisory summary" and granted_record.get("scanStatus") == "COMPLETED"
          and "department" in granted_record, f"got {granted_record}")
    check("active grant -> A's record and latest vitals visible",
          record_a in [r["id"] for r in overview["medicalRecords"]]
          and next(p for p in overview["patients"] if p["id"] == patient_a)["latestVitals"] is not None)
    check("grant is per doctor -> doctor 2 does not see A", patient_a not in overview_patient_ids(tok_doctor_2))

    response = call("POST", "/api/backend/doctor/access/end", tok_doctor_2, {"patientId": patient_a, "doctorId": doctor})
    check("doctor 2 calling access/end (doctorId spoofed) -> 200", response.status_code == 200, f"got {response.status_code}")
    check("doctor 2 cannot end doctor 1's grant -> A still listed", patient_a in overview_patient_ids(tok_doctor))

    response = call("POST", "/api/backend/doctor/access/end", tok_doctor, {"patientId": patient_a})
    check("doctor ends access -> 200", response.status_code == 200, f"got {response.status_code}")
    check("after access/end -> patient A not listed", patient_a not in overview_patient_ids(tok_doctor))

    print("\n[7] OTP unlock flow (patient-issued OTP -> doctor unlock -> end)")
    response = call("POST", "/api/backend/patient/dashboard-otp", tok_a, {"userId": patient_b, "user_id": patient_b})
    otp = (response.json().get("otp") or {}).get("code")
    check("A generates dashboard OTP -> 200", response.status_code == 200 and bool(otp), f"got {response.status_code}")
    check(
        "OTP belongs to A (body userId=B ignored)",
        scalar("select patient_id from patient_dashboard_otps where code = :c", c=otp) == patient_a,
    )
    response = call("POST", "/api/backend/doctor/access/unlock", tok_doctor, {"patientAccessOtp": otp, "doctorId": doctor_2})
    check("doctor unlocks with OTP -> 200", response.status_code == 200, f"got {response.status_code}")
    check("after unlock -> A listed for doctor", patient_a in overview_patient_ids(tok_doctor))
    check("after unlock -> A not listed for doctor 2 (doctorId spoof ignored)", patient_a not in overview_patient_ids(tok_doctor_2))
    response = call("POST", "/api/backend/doctor/access/unlock", tok_doctor, {"patientAccessOtp": otp})
    check("OTP is single-use -> second unlock 404", response.status_code == 404, f"got {response.status_code}")
    call("POST", "/api/backend/doctor/access/end", tok_doctor, {"patientId": patient_a})
    check("after end -> A not listed", patient_a not in overview_patient_ids(tok_doctor))

    otp = call("POST", "/api/backend/patient/dashboard-otp", tok_a).json()["otp"]["code"]
    response = call("POST", "/api/backend/appointments", tok_a,
                    {"doctor_id": doctor_2, "date": next_slot(doctor_2), "reason": "Second opinion", "patient_access_otp": otp})
    check("booking with an OTP -> access granted", response.status_code == 201 and response.json().get("dashboard_access_granted") is True,
          f"got {response.status_code}")
    check("booked doctor (2) now sees A", patient_a in overview_patient_ids(tok_doctor_2))
    check("other doctor still does not", patient_a not in overview_patient_ids(tok_doctor))
    call("POST", "/api/backend/doctor/access/end", tok_doctor_2, {"patientId": patient_a})

    print("\n[8] OTP brute-force limit: 5 failures per 15 minutes")
    tok_bf = token_for(doctor_bf)
    for attempt in range(1, 6):
        response = call("POST", "/api/backend/doctor/access/unlock", tok_bf, {"patientAccessOtp": f"wrong-{attempt}"})
        check(f"wrong unlock attempt {attempt} -> 404", response.status_code == 404, f"got {response.status_code}")
    response = call("POST", "/api/backend/doctor/access/unlock", tok_bf, {"patientAccessOtp": "wrong-6"})
    check("6th attempt -> 429 with Retry-After and a clear message",
          response.status_code == 429 and int(response.headers.get("retry-after", "0")) > 0
          and "Too many incorrect OTP attempts" in response.json().get("detail", ""),
          f"got {response.status_code} {response.headers.get('retry-after')}")
    good_otp = call("POST", "/api/backend/patient/dashboard-otp", tok_a).json()["otp"]["code"]
    response = call("POST", "/api/backend/doctor/access/unlock", tok_bf, {"patientAccessOtp": good_otp})
    check("while blocked, even a correct OTP -> 429 (and the OTP stays unused)", response.status_code == 429,
          f"got {response.status_code}")
    check("blocked doctor does not see A", patient_a not in overview_patient_ids(tok_bf))
    check("5 failed unlock attempts logged for that doctor (code not stored)",
          scalar("select count(*) from dashboard_otp_attempts where actor_id = :id and context = 'unlock' "
                 "and succeeded = false", id=doctor_bf) == 5)
    response = call("POST", "/api/backend/doctor/access/unlock", tok_doctor_2, {"patientAccessOtp": good_otp})
    check("another doctor is not affected -> unlock 200", response.status_code == 200, f"got {response.status_code}")
    call("POST", "/api/backend/doctor/access/end", tok_doctor_2, {"patientId": patient_a})

    appointments_before = scalar("select count(*) from appointments where user_id = :id", id=patient_b)
    for attempt in range(1, 6):
        response = call("POST", "/api/backend/appointments", tok_b,
                        {"doctor_id": doctor, "date": next_slot(doctor), "patient_access_otp": f"wrong-{attempt}"})
        check(f"booking with wrong OTP {attempt} -> 201, no access shared",
              response.status_code == 201 and response.json().get("dashboard_access_granted") is False,
              f"got {response.status_code}")
    response = call("POST", "/api/backend/appointments", tok_b, {"doctor_id": doctor, "date": next_slot(doctor), "patient_access_otp": "wrong-6"})
    check("6th booking with an OTP -> 429", response.status_code == 429, f"got {response.status_code}")
    check("the blocked booking created no appointment",
          scalar("select count(*) from appointments where user_id = :id", id=patient_b) == appointments_before + 5)
    response = call("POST", "/api/backend/appointments", tok_b, {"doctor_id": doctor, "date": next_slot(doctor)})
    check("booking without an OTP still works while blocked -> 201", response.status_code == 201, f"got {response.status_code}")

    print("\n[9] Uploaded files: auth and ownership or an active grant")
    file_bytes = b"\x89PNG\r\n\x1a\n phase0-test file " + secrets.token_bytes(8)
    relative = f"{patient_a}/general/phase0-{secrets.token_hex(8)}.png"
    disk_file = get_upload_root_path() / relative
    disk_file.parent.mkdir(parents=True, exist_ok=True)
    disk_file.write_bytes(file_bytes)
    file_report = create_record(patient_a, f"{TEST_PREFIX}file report", file_url=f"uploads/{relative}")
    orphan = f"{patient_a}/general/phase0-orphan-{secrets.token_hex(8)}.png"
    (get_upload_root_path() / orphan).write_bytes(b"orphan")
    files_path = f"/api/backend/files/{relative}"

    response = call("GET", f"/api/backend/uploads/{relative}")
    check("old public /uploads URL is gone -> 404", response.status_code == 404, f"got {response.status_code}")
    check("no auth -> 401", call("GET", files_path).status_code == 401)
    response = call("GET", files_path, tok_b)
    check("another patient -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("GET", files_path, tok_a)
    check("owner -> 200 with the exact file", response.status_code == 200 and response.content == file_bytes,
          f"got {response.status_code}")
    response = call("GET", files_path, tok_doctor)
    check("doctor without a grant -> 403", response.status_code == 403, f"got {response.status_code}")
    create_grant(doctor, patient_a, expires_in=timedelta(hours=1))
    response = call("GET", files_path, tok_doctor)
    check("doctor with an active grant -> 200", response.status_code == 200, f"got {response.status_code}")
    doctor_expiry = int(time.time()) + 300
    doctor_link = (f"{files_path}?exp={doctor_expiry}&uid={doctor}"
                   f"&sig={_signature(relative, doctor_expiry, doctor)}")
    check("doctor's signed link works during the grant", call("GET", doctor_link).status_code == 200)
    call("POST", "/api/backend/doctor/access/end", tok_doctor, {"patientId": patient_a})
    check("after access/end -> doctor gets 403 by token", call("GET", files_path, tok_doctor).status_code == 403)
    check("after access/end -> doctor's unexpired signed link also stops (403)", call("GET", doctor_link).status_code == 403)

    detail = call("GET", f"/api/backend/patient/reports/{file_report}", tok_a).json()
    link = detail.get("file_access_url") or ""
    check("the report gives the owner a signed link, not the raw path",
          link.startswith(f"files/{relative}?") and "sig=" in link and "uploads/" not in link, f"got {link[:80]!r}")
    response = call("GET", f"/api/backend/{link}")
    check("signed link works without a token -> 200", response.status_code == 200 and response.content == file_bytes,
          f"got {response.status_code}")
    check("tampered signature -> 403", call("GET", f"/api/backend/{link[:-4]}0000").status_code == 403)
    check("same signature for another viewer (uid=B) -> 403",
          call("GET", f"/api/backend/{link.replace(f'uid={patient_a}', f'uid={patient_b}')}").status_code == 403)
    past = int(time.time()) - 5
    expired = f"{files_path}?exp={past}&uid={patient_a}&sig={_signature(relative, past, patient_a)}"
    check("expired link -> 403", call("GET", expired).status_code == 403)
    response = call("GET", f"/api/backend/files/{orphan}", tok_a)
    check("file no report refers to -> 404 even for its owner", response.status_code == 404, f"got {response.status_code}")
    response = call("GET", f"/api/backend/files/{patient_a}/%2E%2E/%2E%2E/main.py", tok_a)
    check("path traversal is not served", response.status_code in (400, 404), f"got {response.status_code}")

    print("\n[13] Pagination: lists return { items, total, limit, offset }")
    for i in range(3):
        create_record(patient_a, f"{TEST_PREFIX}page {i}")
    a_total = scalar("select count(*) from patient_reports where patient_id = :id", id=patient_a)
    first = call("GET", "/api/backend/patient/reports", tok_a).json()
    check("reports list is a page envelope", set(first) == {"items", "total", "limit", "offset"}, f"got keys {sorted(first)}")
    check("default page: limit 20, offset 0, total = A's reports",
          first["limit"] == 20 and first["offset"] == 0 and first["total"] == a_total, f"got {first['limit']}/{first['offset']}/{first['total']}")
    page1 = call("GET", "/api/backend/patient/reports?limit=2&offset=0", tok_a).json()
    page2 = call("GET", "/api/backend/patient/reports?limit=2&offset=2", tok_a).json()
    check("limit=2 returns 2 items with the full total", len(page1["items"]) == 2 and page1["total"] == a_total)
    check("consecutive pages do not overlap",
          not {r["id"] for r in page1["items"]} & {r["id"] for r in page2["items"]})
    past_end = call("GET", f"/api/backend/patient/reports?limit=5&offset={a_total}", tok_a).json()
    check("a page past the end is empty and keeps the total", past_end["items"] == [] and past_end["total"] == a_total)
    b_page = call("GET", "/api/backend/patient/reports", tok_b).json()
    check("B's total counts only B's reports",
          b_page["total"] == scalar("select count(*) from patient_reports where patient_id = :id", id=patient_b))
    for bad in ("limit=0", "limit=101", "offset=-1", "limit=abc"):
        response = call("GET", f"/api/backend/patient/reports?{bad}", tok_a)
        check(f"?{bad} -> 422", response.status_code == 422, f"got {response.status_code}")
    users_page = call("GET", "/api/backend/api/admin/users?limit=1", tok_admin).json()
    check("admin users: limit=1 returns one user and the total user count",
          len(users_page["items"]) == 1 and users_page["total"] == scalar("select count(*) from users"))
    stats = call("GET", "/api/backend/api/admin/dashboard-stats", tok_admin).json()
    check("admin stats include patient and active-user counts",
          stats.get("totalPatients") == scalar("select count(*) from users where lower(role) = 'patient'")
          and stats.get("activeUsers") == scalar("select count(*) from users where is_active"), f"got {stats}")

    print("\n[13b] Report search and filters (server-side)")
    tiles = call("GET", "/api/backend/patient/reports/tiles", tok_a).json()
    check("department tiles are exactly the six departments lib/departments.ts mirrors",
          {t["slug"] for t in tiles} == {"heart", "kidney", "ortho", "general", "medicine", "lab"},
          f"got {sorted(t['slug'] for t in tiles)}")
    heart_report = create_record(patient_a, f"{TEST_PREFIX}searchable heart scan")
    with engine.begin() as conn:
        conn.execute(text("update patient_reports set validation_status = 'HIGH_RISK', "
                          "department_id = (select id from medical_departments where name = 'Heart') where id = :id"),
                     {"id": heart_report})
    found = call("GET", "/api/backend/patient/reports?q=SEARCHABLE%20heart", tok_a).json()
    check("q matches the title, case-insensitively", [r["id"] for r in found["items"]] == [heart_report] and found["total"] == 1,
          f"got {[r['id'] for r in found['items']]}")
    check("q does not search other patients' reports",
          call("GET", "/api/backend/patient/reports?q=searchable", tok_b).json()["total"] == 0)
    wildcard = call("GET", "/api/backend/patient/reports?q=%25", tok_a).json()
    check("% in q is matched literally, not as a wildcard", wildcard["total"] == 0, f"got total {wildcard['total']}")
    heart = call("GET", "/api/backend/patient/reports?department=heart", tok_a).json()
    check("department=heart returns only heart reports",
          heart["total"] >= 1 and all(r["department_slug"] == "heart" for r in heart["items"]), f"got {heart['total']}")
    risky = call("GET", "/api/backend/patient/reports?validation_status=HIGH_RISK", tok_a).json()
    check("validation_status=HIGH_RISK returns only high-risk reports",
          heart_report in [r["id"] for r in risky["items"]] and all(r["validation_status"] == "HIGH_RISK" for r in risky["items"]))
    for bad in ("department=liver", "validation_status=STABLE", "q=" + "x" * 101):
        response = call("GET", f"/api/backend/patient/reports?{bad}", tok_a)
        check(f"?{bad[:30]} -> 422", response.status_code == 422, f"got {response.status_code}")

    print("\n[13c] Account and doctor-review responses")
    me = call("GET", "/api/backend/api/auth/me", tok_a)
    check("/auth/me -> 200 for an address on a reserved domain (@example.test)",
          me.status_code == 200 and me.json().get("email", "").endswith("@example.test"), f"got {me.status_code}")
    reviewed = create_user("doctor-review", "DOCTOR", doctor_status="pending")
    profile_id = scalar("select id from doctor_profiles where user_id = :uid", uid=reviewed)
    # Only send the user id when it matches no profile, so this can never touch a real doctor.
    if scalar("select count(*) from doctor_profiles where id = :id", id=reviewed) == 0:
        response = call("PATCH", f"/api/backend/api/admin/doctors/{reviewed}/approve", tok_admin)
        check("approve with the doctor's user id (not the profile id) -> 404", response.status_code == 404,
              f"got {response.status_code}")
    response = call("PATCH", f"/api/backend/api/admin/doctors/{profile_id}/approve", tok_admin)
    check("approve with the profile id -> 200 and approved",
          response.status_code == 200 and response.json().get("verificationStatus") == "approved"
          and scalar("select is_identity_verified from users where id = :id", id=reviewed) is True)
    response = call("PATCH", f"/api/backend/api/admin/doctors/{profile_id}/reject", tok_admin)
    check("reject with the profile id -> 200 and rejected",
          response.status_code == 200
          and scalar("select verification_status from doctor_profiles where id = :id", id=profile_id) == "rejected")

    taken = f"{8_000_000 + secrets.randbelow(900_000)}-P"
    with engine.begin() as conn:
        conn.execute(text("update doctor_profiles set pmdc_number = :p where user_id = :id"), {"p": taken, "id": doctor})
    response = call("POST", "/api/auth/signup", body={
        "role": "doctor", "full_name": f"{TEST_PREFIX}pmdc duplicate", "email": f"{TEST_PREFIX}pmdc-dup-{run_id}@example.com",
        "password": "Pmdc-test-pass-1", "confirm_password": "Pmdc-test-pass-1", "phone": f"0311 {secrets.randbelow(10**7):07d}",
        "pmdc_number": f"  {taken.upper()} ", "specialization": "Cardiology", "hospital_name": "Test Hospital",
    })
    check("doctor signup with a PMDC number already in use (other case and spacing) -> 409",
          response.status_code == 409 and "PMDC" in response.json().get("detail", ""), f"got {response.status_code}")
    twin = create_user("doctor-pmdc-twin", "DOCTOR", doctor_status="pending")
    with engine.begin() as conn:
        conn.execute(text("update doctor_profiles set pmdc_number = :p where user_id = :id"), {"p": taken, "id": twin})
    twin_profile = scalar("select id from doctor_profiles where user_id = :uid", uid=twin)
    response = call("PATCH", f"/api/backend/api/admin/doctors/{twin_profile}/approve", tok_admin)
    check("approving a doctor whose PMDC number an approved doctor already has -> 409, stays pending",
          response.status_code == 409
          and scalar("select verification_status from doctor_profiles where id = :id", id=twin_profile) == "pending",
          f"got {response.status_code}")
    response = call("PATCH", f"/api/backend/api/admin/doctors/{twin_profile}/reject", tok_admin)
    check("that doctor can still be rejected -> 200", response.status_code == 200, f"got {response.status_code}")


    print("\n[14] Rate limits")
    limited_patient = create_user("ratelimit", "PATIENT")
    tok_limited = token_for(limited_patient)
    responses = [call("POST", "/api/backend/patient/dashboard-otp", tok_limited) for _ in range(11)]
    statuses = [r.status_code for r in responses]
    check("10 access codes in a minute are allowed", statuses[:10] == [200] * 10, f"got {statuses}")
    blocked = responses[10]
    check("the 11th -> 429 with Retry-After and a readable detail",
          blocked.status_code == 429 and blocked.headers.get("retry-after", "").isdigit()
          and "Too many requests" in blocked.json().get("detail", ""), f"got {blocked.status_code} {blocked.text[:120]}")
    response = call("POST", "/api/backend/patient/dashboard-otp", tok_a)
    check("limits are per user: another patient is not affected", response.status_code == 200, f"got {response.status_code}")

    email = f"{TEST_PREFIX}nobody-{secrets.token_hex(4)}@example.com"
    attempts = [call("POST", "/api/auth/login", body={"email": email, "password": "wrong-password"}) for _ in range(6)]
    codes = [r.status_code for r in attempts]
    check("5 failed logins for one email -> 401 each", codes[:5] == [401] * 5, f"got {codes}")
    check("6th login attempt for that email -> 429 with Retry-After",
          codes[5] == 429 and attempts[5].headers.get("retry-after", "").isdigit(), f"got {codes}")
    other = call("POST", "/api/auth/login", body={"email": f"{TEST_PREFIX}other-{secrets.token_hex(4)}@example.com",
                                                   "password": "wrong-password"})
    check("another email from the same client is still allowed (401, not 429)", other.status_code == 401,
          f"got {other.status_code}")

    print("\n[15] Clinical summary for doctors: active grant only, versions, out of date after changes")
    patient_c = create_user("patient-summary", "PATIENT")
    doctor_c = create_user("doctor-summary", "DOCTOR")
    tok_c, tok_doctor_c = token_for(patient_c), token_for(doctor_c)
    report_c = create_record(patient_c, f"{TEST_PREFIX}summary report")
    base = f"/api/backend/doctor/patients/{patient_c}/clinical-summary"
    for method, path in (("GET", base), ("POST", f"{base}/regenerate"), ("GET", f"{base}/versions")):
        response = call(method, path, tok_doctor_c)
        check(f"doctor without a grant: {method} {path.replace(base, '...')} -> 403", response.status_code == 403,
              f"got {response.status_code}")
    check("patient reading their own clinical summary -> 403", call("GET", base, tok_c).status_code == 403)
    check("admin -> 403", call("GET", base, tok_admin).status_code == 403)
    response = call("GET", f"/api/backend/doctor/patients/{doctor_2}/clinical-summary", tok_doctor_c)
    check("a non-patient id -> 403 (same answer, nothing revealed)", response.status_code == 403, f"got {response.status_code}")

    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=1))
    if configured_paid_keys():
        # Regenerating asks the AI to word "Key findings" when a key is set: tests never call paid services.
        print("  SKIP  clinical summary regeneration: an AI key is configured")
    else:
        before = scalar("select md5(row(r.*)::text) from patient_reports r where id = :id", id=report_c)
        response = call("GET", base, tok_doctor_c)
        body = response.json() if response.status_code == 200 else {}
        check("with a grant: no summary yet, out of date", response.status_code == 200 and body.get("summary") is None
              and body.get("out_of_date") is True, f"got {response.status_code} {str(body)[:160]}")
        response = call("POST", f"{base}/regenerate", tok_doctor_c)
        body = response.json() if response.status_code == 200 else {}
        summary = body.get("summary") or {}
        content = summary.get("content") or {}
        check("regenerate -> version 1 created", response.status_code == 200 and body.get("created") is True
              and summary.get("version") == 1, f"got {response.status_code} {str(body)[:160]}")
        abnormal = content.get("abnormal_results") or []
        check("the low hemoglobin is an abnormal result, linked to its report",
              any(f.get("subject") == "Hemoglobin" and any(s.get("report_id") == report_c for s in f.get("sources", []))
                  for f in abnormal), f"abnormal={str(abnormal)[:200]}")
        check("a report with no stored entities was re-read for the summary",
              (content.get("inputs") or {}).get("reports_reread") == 1, f"inputs={content.get('inputs', {}).get('reports_reread')}")
        check("the rule-based note is attached and no 'safe'/'normal' wording",
              "fixed rules" in summary.get("note", "") and " safe" not in summary.get("summary_text", "").lower()
              and "normal" not in summary.get("summary_text", "").lower().replace("abnormal", ""))
        check("building the summary did not change the report row",
              scalar("select md5(row(r.*)::text) from patient_reports r where id = :id", id=report_c) == before)
        response = call("POST", f"{base}/regenerate", tok_doctor_c)
        check("regenerate with nothing changed -> no new version",
              response.status_code == 200 and response.json().get("created") is False
              and response.json()["summary"]["version"] == 1, f"got {response.status_code}")
        check("current summary is up to date", call("GET", base, tok_doctor_c).json().get("out_of_date") is False)

        create_record(patient_c, f"{TEST_PREFIX}summary report 2")
        body = call("GET", base, tok_doctor_c).json()
        check("after a new report -> out of date, saying what changed", body.get("out_of_date") is True
              and "1 report(s) added" in (body.get("pending_change") or "") and body["summary"]["version"] == 1,
              f"got {str(body)[:200]}")
        response = call("POST", f"{base}/regenerate", tok_doctor_c)
        check("regenerate -> version 2 with its trigger", response.status_code == 200
              and response.json()["summary"]["version"] == 2 and "report(s) added" in response.json()["summary"]["trigger"],
              f"got {response.status_code} {response.text[:160]}")
        response = call("GET", f"{base}/versions", tok_doctor_c)
        page = response.json() if response.status_code == 200 else {}
        check("both versions are kept, newest first, only the newest current",
              page.get("total") == 2 and [v["version"] for v in page.get("items", [])] == [2, 1]
              and [v["is_current"] for v in page.get("items", [])] == [True, False], f"got {str(page)[:200]}")
        check("an earlier version can be read", call("GET", f"{base}/versions/1", tok_doctor_c).status_code == 200)
        check("a version that does not exist -> 404", call("GET", f"{base}/versions/99", tok_doctor_c).status_code == 404)
        check("another doctor (no grant) cannot read a version -> 403",
              call("GET", f"{base}/versions/1", tok_doctor_2).status_code == 403)
    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=-1))
    check("after the grant expires -> 403", call("GET", base, tok_doctor_c).status_code == 403)

    print("\n[16] Everyday wording for patients")
    response = call("GET", f"/api/backend/patient/reports/{report_c}", tok_c)
    body = response.json() if response.status_code == 200 else {}
    check("a patient's report carries conditions_mentioned (a list)",
          response.status_code == 200 and isinstance(body.get("conditions_mentioned"), list), f"got {response.status_code}")
    check("every rule alert on it has a patient_message",
          all(w.get("patient_message") for w in body.get("rule_warnings", [])), f"warnings={body.get('rule_warnings')}")
    with engine.begin() as conn:
        conn.execute(text(
            "update patient_reports set validation_status = 'HIGH_RISK', rule_warnings = :w where id = :id"),
            {"id": report_c, "w": json.dumps([{"code": "KIDNEY_RISK_EGFR", "severity": "HIGH_RISK",
                                                "message": "eGFR 25 is under 30.", "test": "eGFR",
                                                "value_text": "25", "direction": "low"}])})
    response = call("GET", "/api/backend/patient/health-flags", tok_c)
    flags = response.json() if response.status_code == 200 else {}
    check("health flags: the patient's own red flag, in everyday words, from the rules",
          response.status_code == 200 and flags.get("high_risk_reports") == 1
          and "kidney filtering rate" in " ".join(flags["reports"][0]["messages"]), f"got {response.status_code} {str(flags)[:200]}")
    response = call("GET", "/api/backend/patient/health-flags", tok_a)
    check("another patient does not see those flags",
          response.status_code == 200 and all(r["report_id"] != report_c for r in response.json()["reports"]))

    print("\n[17] Medication safety and prescribing (MSCIE): grant only, alerts, justifications, patient view")
    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=1))
    with engine.begin() as conn:
        conn.execute(text(
            "update patient_reports set formatted_text = :t, extracted_text = :t, structured_data_json = '{}' "
            "where id = :id"), {"id": report_c, "t": "Date: 01/03/2026\nK/C/O CKD.\nTab Warfarin 5 mg OD\neGFR 25 mL/min/1.73m2"})
    rx_base = f"/api/backend/doctor/patients/{patient_c}"
    response = call("GET", "/api/backend/doctor/medicines/search?q=glucophag", tok_doctor_c)
    check("medicine search finds a misspelled brand (fuzzy)",
          response.status_code == 200 and response.json()[0]["generic"] == "metformin", f"got {response.status_code}")
    body = {"items": [{"medicine": "Brufen", "dose": "400 mg", "frequency": "TDS"}]}
    for who, token in (("doctor without a grant", tok_doctor_2), ("the patient", tok_c), ("an admin", tok_admin)):
        response = call("POST", f"{rx_base}/medication-check", token, body)
        check(f"medication check by {who} -> 403", response.status_code == 403, f"got {response.status_code}")
    response = call("POST", f"{rx_base}/medication-check", tok_doctor_c, body)
    checked = response.json() if response.status_code == 200 else {}
    item = (checked.get("items") or [{}])[0]
    titles = [a["title"] for a in item.get("alerts", [])]
    check("check: ibuprofen with warfarin, CKD and eGFR 25 are high-risk alerts, with evidence and sources",
          response.status_code == 200 and checked.get("worst") == "major"
          and any("warfarin" in t for t in titles) and any("chronic kidney disease" in t for t in titles)
          and all(a["sources"] or a["kind"] in ("duplicate", "previous_response") for a in item.get("alerts", [])),
          f"got {response.status_code} {titles}")
    check("check: paracetamol is the first ranked alternative, with its score breakdown",
          (item.get("alternatives") or [{}])[0].get("medicine") == "paracetamol"
          and len(item["alternatives"][0].get("breakdown", [])) == 4, f"alternatives={item.get('alternatives')}")
    check("check: the context shows the record used (warfarin, CKD, eGFR 25)",
          any(m["name"] == "warfarin" for m in checked.get("context", {}).get("medicines", []))
          and any(lab["test"] == "eGFR" and lab["value"] == 25 for lab in checked.get("context", {}).get("labs", [])))
    response = call("POST", f"{rx_base}/prescriptions", tok_doctor_c, body)
    check("confirming without justifications -> 422 listing each high-risk alert",
          response.status_code == 422 and len(response.json()["detail"]["alerts"]) == len(checked.get("requires_justification", [])),
          f"got {response.status_code}")
    short = [{"alert_id": a["id"], "justification": "too short"} for a in checked.get("requires_justification", [])]
    response = call("POST", f"{rx_base}/prescriptions", tok_doctor_c, {**body, "overrides": short})
    check("a justification under 20 characters -> 422", response.status_code == 422, f"got {response.status_code}")
    overrides = [{"alert_id": a["id"], "justification": "Short course agreed with nephrology; INR checked in 3 days."}
                 for a in checked.get("requires_justification", [])]
    response = call("POST", f"{rx_base}/prescriptions", tok_doctor_c, {**body, "overrides": overrides, "notes": "phase0-test"})
    created = response.json() if response.status_code == 201 else {}
    check("confirming with a justification for each high-risk alert -> 201, overrides stored",
          response.status_code == 201 and len(created.get("overrides", [])) == len(overrides)
          and scalar("select count(*) from safety_overrides where prescription_id = :id", id=created.get("id") or 0) == len(overrides),
          f"got {response.status_code} {response.text[:160]}")
    response = call("GET", "/api/backend/patient/prescriptions", tok_c)
    mine = response.json() if response.status_code == 200 else {}
    check("the patient sees it read-only, in everyday words (three times a day)",
          response.status_code == 200 and mine.get("total") == 1
          and "three times a day" in mine["items"][0]["items"][0]["how_to_take"]
          and "alerts" not in mine["items"][0], f"got {response.status_code} {str(mine)[:160]}")
    check("another patient sees none of it", call("GET", "/api/backend/patient/prescriptions", tok_a).json().get("total") == 0)
    cancel = f"{rx_base}/prescriptions/{created.get('id')}/cancel"
    check("another doctor without a grant cannot cancel -> 403", call("POST", cancel, tok_doctor_2).status_code == 403)
    response = call("POST", cancel, tok_doctor_c)
    check("the prescribing doctor cancels -> cancelled", response.status_code == 200 and response.json()["status"] == "cancelled")
    check("cancelling twice -> 409", call("POST", cancel, tok_doctor_c).status_code == 409)
    response = call("GET", f"{rx_base}/prescriptions", tok_doctor_c)
    check("the doctor's list keeps the cancelled prescription", response.status_code == 200 and response.json()["total"] == 1)
    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=-1))
    check("after the grant expires -> 403", call("POST", f"{rx_base}/medication-check", tok_doctor_c, body).status_code == 403)

    print("\n[17b] Health record (Module 2) and the audit log")
    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=1))
    hr = "/api/backend/patient/health-record"
    dhr = f"/api/backend/doctor/patients/{patient_c}/health-record"
    check("health record without a token -> 401", call("GET", hr).status_code == 401)
    check("a doctor on the patient's health record -> 403", call("GET", hr, tok_doctor_c).status_code == 403)
    check("a patient on the doctor's health record -> 403", call("GET", dhr, tok_c).status_code == 403)
    check("the options list needs sign-in", call("GET", "/api/backend/health-record/options").status_code == 401)
    options = call("GET", "/api/backend/health-record/options", tok_c).json()
    check("options offer known allergy and condition names", "penicillin" in options["allergies"] and "diabetes mellitus" in options["conditions"])

    response = call("POST", hr, tok_c, {"kind": "allergy", "name": "  penicillin ", "severity": "severe", "note": "rash"})
    allergy = response.json() if response.status_code == 201 else {}
    check("C adds an allergy -> 201, trimmed, recognised by the medication check",
          response.status_code == 201 and allergy.get("name") == "penicillin" and allergy.get("recognised") is True
          and allergy.get("added_by_role") == "patient", f"got {response.status_code}")
    check("the same allergy twice -> 409", call("POST", hr, tok_c, {"kind": "allergy", "name": "Penicillin"}).status_code == 409)
    check("an unknown kind -> 422", call("POST", hr, tok_c, {"kind": "medicine", "name": "aspirin"}).status_code == 422)
    check("a date in the future -> 422", call("POST", hr, tok_c, {"kind": "diagnosis", "name": "asthma", "since": "2999-01-01"}).status_code == 422)
    response = call("POST", hr, tok_c, {"kind": "diagnosis", "name": "feeling sad on mondays"})
    check("a word the check does not know is saved but marked not recognised",
          response.status_code == 201 and response.json().get("recognised") is False)
    unknown_id = response.json().get("id")

    entry_path = f"{hr}/{allergy.get('id')}"
    check("A cannot change C's entry -> 404", call("PATCH", entry_path, tok_a, {"note": "x"}).status_code == 404)
    check("A cannot remove C's entry -> 404", call("DELETE", entry_path, tok_a).status_code == 404)
    check("A's list never shows C's entries", all(e["id"] != allergy.get("id") for e in call("GET", hr, tok_a).json()))

    check("a doctor without a grant cannot read the list -> 403", call("GET", dhr, tok_doctor_2).status_code == 403)
    check("or add to it -> 403", call("POST", dhr, tok_doctor_2, {"kind": "diagnosis", "name": "asthma"}).status_code == 403)
    check("or change an entry -> 403", call("PATCH", f"{dhr}/{allergy.get('id')}", tok_doctor_2, {"note": "x"}).status_code == 403)
    listed = call("GET", dhr, tok_doctor_c)
    check("the granted doctor reads the list", listed.status_code == 200 and any(e["name"] == "penicillin" for e in listed.json()))
    call("GET", dhr, tok_doctor_c)
    check("two views inside the window leave one audit row",
          scalar("select count(*) from audit_logs where actor_id = :d and action = 'record.view' and detail = 'health record'", d=doctor_c) == 1)

    response = call("POST", dhr, tok_doctor_c, {"kind": "chronic_condition", "name": "diabetes"})
    added = response.json() if response.status_code == 201 else {}
    check("the granted doctor adds a condition, marked as added by a doctor, with their name",
          response.status_code == 201 and added.get("added_by_role") == "doctor" and added.get("added_by_name"), f"got {response.status_code}")
    check("the doctor changes an entry", call("PATCH", f"{dhr}/{added.get('id')}", tok_doctor_c, {"note": "on tablets"}).json().get("note") == "on tablets")
    check("a doctor cannot remove entries (no such route) -> 405", call("DELETE", f"{dhr}/{added.get('id')}", tok_doctor_c).status_code == 405)
    check("the patient sees who added it", any(e["id"] == added.get("id") and e["added_by_role"] == "doctor" for e in call("GET", hr, tok_c).json()))
    check("severity on a condition is refused -> 422", call("PATCH", f"{dhr}/{added.get('id')}", tok_doctor_c, {"severity": "mild"}).status_code == 422)

    rx = call("POST", f"/api/backend/doctor/patients/{patient_c}/medication-check", tok_doctor_c,
              {"items": [{"medicine": "amoxicillin", "dose": "500 mg", "frequency": "TDS"}]})
    checked = rx.json() if rx.status_code == 200 else {}
    alerts = [a for it in checked.get("items", []) for a in it.get("alerts", [])]
    check("the medication check uses the recorded penicillin allergy", rx.status_code == 200 and any(a["kind"] == "allergy" for a in alerts),
          f"got {rx.status_code} {[a['title'] for a in alerts]}")
    check("and shows where it came from", any("health record" in a.get("source", "") for a in checked.get("context", {}).get("allergies", [])),
          f"got {checked.get('context', {}).get('allergies')}")
    check("the patient resolves the allergy", call("PATCH", entry_path, tok_c, {"status": "resolved"}).json().get("status") == "resolved")
    rx = call("POST", f"/api/backend/doctor/patients/{patient_c}/medication-check", tok_doctor_c,
              {"items": [{"medicine": "amoxicillin", "dose": "500 mg", "frequency": "TDS"}]})
    check("a resolved allergy is no longer used by the check",
          not any(a["kind"] == "allergy" for it in rx.json().get("items", []) for a in it.get("alerts", [])))
    response = call("DELETE", f"{hr}/{unknown_id}", tok_c)
    check("the patient removes an entry -> 200 and it leaves the list",
          response.status_code == 200 and all(e["id"] != unknown_id for e in call("GET", hr, tok_c).json()))
    check("a removed entry cannot be changed -> 404", call("PATCH", f"{hr}/{unknown_id}", tok_c, {"note": "x"}).status_code == 404)
    check("the row is kept (hidden, not deleted)", scalar("select status from patient_health_entries where id = :i", i=unknown_id) == "removed")

    audit_path = "/api/backend/api/admin/audit-log"
    check("audit log without a token -> 401", call("GET", audit_path).status_code == 401)
    check("a patient on the audit log -> 403", call("GET", audit_path, tok_c).status_code == 403)
    check("a doctor on the audit log -> 403", call("GET", audit_path, tok_doctor_c).status_code == 403)
    check("an unknown action filter -> 422", call("GET", audit_path + "?action=bogus", tok_admin).status_code == 422)
    response = call("GET", f"{audit_path}?action=health_entry.add&subject={patient_c}", tok_admin)
    rows = response.json() if response.status_code == 200 else {}
    check("admin filters the log by action and patient, with names and a total",
          response.status_code == 200 and rows.get("total", 0) >= 3 and all(r["action"] == "health_entry.add" and r["patient_id"] == patient_c for r in rows["items"])
          and all(r["actor_name"] for r in rows["items"]), f"got {response.status_code}")
    check("the log filters by who acted", all(r["actor_role"] == "doctor" for r in call("GET", f"{audit_path}?actor_role=doctor&subject={patient_c}", tok_admin).json()["items"]))
    check("the log never holds what was entered (no allergy or condition names)",
          "penicillin" not in call("GET", f"{audit_path}?subject={patient_c}&limit=100", tok_admin).text.lower())
    check("an admin can read the action list", "record.view" in call("GET", audit_path + "/actions", tok_admin).json())
    check("prescribing is on the log (confirm and cancel)",
          scalar("select count(*) from audit_logs where patient_id = :p and action = 'prescription.confirm'", p=patient_c) >= 1
          and scalar("select count(*) from audit_logs where patient_id = :p and action = 'prescription.cancel'", p=patient_c) >= 1)
    check("access grants and one-time codes are on the log",
          scalar("select count(*) from audit_logs where patient_id = :p and action = 'access.grant'", p=patient_a) >= 1
          and scalar("select count(*) from audit_logs where patient_id = :p and action = 'access.code_issued'", p=patient_a) >= 1)

    history = call("GET", "/api/backend/patient/access-history", tok_c)
    check("C sees who accessed their record, never their own actions",
          history.status_code == 200 and history.json()["total"] >= 2 and all(r["actor_role"] != "patient" for r in history.json()["items"]),
          f"got {history.status_code}")
    check("access history needs sign-in and a patient", call("GET", "/api/backend/patient/access-history").status_code == 401
          and call("GET", "/api/backend/patient/access-history", tok_doctor_c).status_code == 403)
    check("A's access history counts only rows about A",
          call("GET", "/api/backend/patient/access-history", tok_a).json()["total"]
          == scalar("select count(*) from audit_logs where patient_id = :a and actor_role <> 'patient'", a=patient_a))
    create_grant(doctor_c, patient_c, expires_in=timedelta(hours=-1))

    print("\n[17c] Doctor PMDC verification: the format at signup, the admin's check, the badge")
    signup_path = "/api/auth/signup"
    pm_digits = str(9_000_000 + secrets.randbelow(900_000))  # seven digits, unlikely to be anyone's number

    def doctor_signup(tag: str, number: str):
        return httpx.post(f"{BASE_URL}{signup_path}", json={
            "role": "doctor", "full_name": f"{TEST_PREFIX}pmdc {tag}", "email": f"{TEST_PREFIX}pmdc-{tag}-{run_id}@example.com",
            "password": "Pmdc-test-pass-1", "confirm_password": "Pmdc-test-pass-1", "phone": f"0312 {secrets.randbelow(10**7):07d}",
            "pmdc_number": number, "specialization": "Cardiology", "hospital_name": "Test Hospital"}, timeout=60)

    response = doctor_signup("bad", "PMC-12345")
    check("doctor signup with a number that is not in a recognised form -> 422 with an example",
          response.status_code == 422 and "12345-P" in str(response.json().get("detail")), f"got {response.status_code}")
    check("no account was created by a refused signup",
          scalar("select count(*) from users where email like :e", e=f"{TEST_PREFIX}pmdc-bad-%") == 0)
    response = doctor_signup("ok", f" pmdc  {pm_digits} p ")
    check("a number written with a prefix, spaces and lower case is accepted", response.status_code in (200, 201), f"got {response.status_code}")
    check("and stored in the standard form the register searches on",
          scalar("select pmdc_number from doctor_profiles dp join users u on u.id = dp.user_id where u.email = :e",
                 e=f"{TEST_PREFIX}pmdc-ok-{run_id}@example.com") == f"{pm_digits}-P")
    response = doctor_signup("dup", f"{pm_digits}/p")
    check("the same number in another spelling -> 409 already registered", response.status_code == 409, f"got {response.status_code}")
    new_doctor_profile = scalar("select dp.id from doctor_profiles dp join users u on u.id = dp.user_id where u.email = :e",
                                e=f"{TEST_PREFIX}pmdc-ok-{run_id}@example.com")
    check("a new doctor starts not PMDC verified and waits for approval",
          scalar("select not pmdc_verified and verification_status = 'pending' from doctor_profiles where id = :i", i=new_doctor_profile) is True)

    profile_id = scalar("select id from doctor_profiles where user_id = :u", u=doctor)
    admin_doctors = "/api/backend/api/admin/doctors"
    check("admin doctor list without a token -> 401", call("GET", admin_doctors).status_code == 401)
    check("a patient on the admin doctor list -> 403", call("GET", admin_doctors, tok_a).status_code == 403)
    check("a doctor on the admin doctor list -> 403", call("GET", admin_doctors, tok_doctor).status_code == 403)
    response = call("GET", f"{admin_doctors}?approval=pending&limit=100", tok_admin)
    listing = response.json() if response.status_code == 200 else {}
    mine = [d for d in listing.get("items", []) if d["id"] == new_doctor_profile]
    check("admin lists doctors with their PMDC status (format, not checked yet)",
          response.status_code == 200 and mine and mine[0]["pmdc"]["format_ok"] is True and mine[0]["pmdc"]["verified"] is False
          and mine[0]["pmdc"]["message"] == "Not checked yet." and mine[0]["pmdc"]["searchable"] == f"{pm_digits}-P", f"got {response.status_code}")
    check("an unknown filter value -> 422", call("GET", admin_doctors + "?pmdc=maybe", tok_admin).status_code == 422)
    check("the list filters by name or number", any(d["id"] == new_doctor_profile for d in call("GET", f"{admin_doctors}?q={pm_digits}", tok_admin).json()["items"])
          and all(d["id"] != new_doctor_profile for d in call("GET", f"{admin_doctors}?q=zzzqqq-no-such", tok_admin).json()["items"]))

    with engine.begin() as conn:  # an earlier check gave this doctor a real-looking number: back to a placeholder
        conn.execute(text("update doctor_profiles set pmdc_number = :p where id = :i"), {"p": f"{TEST_PREFIX}doctor-{run_id}", "i": profile_id})
    check_path_ = f"/api/backend/api/admin/doctors/{profile_id}/pmdc-check"
    check("a PMDC check without a token -> 401", call("POST", check_path_).status_code == 401)
    check("a patient cannot run a PMDC check -> 403", call("POST", check_path_, tok_a).status_code == 403)
    check("a doctor cannot run a PMDC check -> 403", call("POST", check_path_, tok_doctor).status_code == 403)
    check("a check on a doctor that does not exist -> 404", call("POST", "/api/backend/api/admin/doctors/999999999/pmdc-check", tok_admin).status_code == 404)
    # The test users' numbers are placeholders, so the register is never asked: the answer is "does not fit".
    response = call("POST", check_path_, tok_admin)
    result = response.json() if response.status_code == 200 else {}
    check("a number that does not fit is reported without asking the register",
          response.status_code == 200 and result.get("outcome") == "format_invalid" and result.get("verified") is False
          and result.get("format_ok") is False, f"got {response.status_code} {result}")
    if settings.pmdc_lookup_enabled:
        print("  SKIP  a check that would ask the PMDC register (tests never call outside services; run with DISABLE_PAID_SERVICES=1)")
    else:
        with engine.begin() as conn:
            conn.execute(text("update doctor_profiles set pmdc_number = '123456-P' where id = :i"), {"i": profile_id})
        response = call("POST", check_path_, tok_admin)
        check("with the lookup switched off a real-looking number -> 503 and nothing changes",
              response.status_code == 503 and scalar("select pmdc_verified from doctor_profiles where id = :i", i=profile_id) is False,
              f"got {response.status_code}")
        with engine.begin() as conn:
            conn.execute(text("update doctor_profiles set pmdc_number = :p where id = :i"),
                         {"p": f"{TEST_PREFIX}doctor-{run_id}", "i": profile_id})

    by_hand = f"/api/backend/api/admin/doctors/{profile_id}/pmdc-verified"
    check("a decision by hand without a token -> 401", call("PUT", by_hand, None, {"verified": True, "note": "checked it myself"}).status_code == 401)
    check("a doctor cannot mark themselves verified -> 403", call("PUT", by_hand, tok_doctor, {"verified": True, "note": "checked it myself"}).status_code == 403)
    check("a decision by hand needs a reason -> 422", call("PUT", by_hand, tok_admin, {"verified": True, "note": "ok"}).status_code == 422)
    response = call("PUT", by_hand, tok_admin, {"verified": True, "note": "Checked on pmdc.pk: the name is spelled differently."})
    manual = response.json() if response.status_code == 200 else {}
    check("admin marks the doctor PMDC verified by hand, with the reason and who decided",
          response.status_code == 200 and manual.get("verified") is True and manual.get("outcome") == "manual"
          and "spelled differently" in manual.get("message", "") and manual.get("checked_by_name"), f"got {response.status_code}")
    check("the decision is on the audit log, without the note",
          scalar("select count(*) from audit_logs where action = 'admin.pmdc_manual' and target_id = :i", i=profile_id) == 1
          and "spelled differently" not in call("GET", "/api/backend/api/admin/audit-log?action=admin.pmdc_manual&limit=100", tok_admin).text)
    public = call("GET", "/api/backend/appointments/doctors").json()
    card = next((d for d in public if d["id"] == doctor), {})
    check("the public doctor card carries the PMDC verified flag and date, never the number",
          card.get("pmdc_verified") is True and card.get("pmdc_checked_at") and "pmdc_number" not in card and "pmdc_note" not in card
          and "pmdc_register_name" not in card)
    only = call("GET", "/api/backend/appointments/doctors?pmdc_verified=true").json()
    check("the PMDC verified filter lists only verified doctors", doctor in [d["id"] for d in only] and all(d["pmdc_verified"] for d in only))
    check("without the filter other doctors are still listed, unverified",
          any(d["pmdc_verified"] is False for d in public) or len(public) == len([d for d in public if d["pmdc_verified"]]))
    check("a doctor sees their own PMDC status on the profile",
          call("GET", "/api/backend/profile", tok_doctor).json()["doctor"]["pmdc_verified"] is True)
    response = call("PUT", "/api/backend/doctor/profile", tok_doctor, {
        "hospital_name": "phase0-test hospital", "consultation_modes": ["in_clinic"], "pmdc_verified": False, "pmdc_number": "1-X"})
    check("a doctor cannot change the flag or the number through their profile",
          response.status_code == 200 and scalar("select pmdc_verified from doctor_profiles where id = :i", i=profile_id) is True
          and scalar("select pmdc_number from doctor_profiles where id = :i", i=profile_id) != "1-X", f"got {response.status_code}")
    check("an unapproved doctor is never listed, verified or not",
          all(d["id"] != new_doctor_profile for d in public) and new_doctor_profile not in [d["id"] for d in public])
    response = call("PUT", by_hand, tok_admin, {"verified": False, "note": "Withdrawn: the register shows a different doctor."})
    check("the administrator can take the badge away", response.status_code == 200 and response.json().get("verified") is False
          and all(d["id"] != doctor for d in call("GET", "/api/backend/appointments/doctors?pmdc_verified=true").json()))
    check("a doctor can be approved without any PMDC check",
          call("PATCH", f"/api/backend/api/admin/doctors/{new_doctor_profile}/approve", tok_admin).status_code == 200
          and scalar("select verification_status = 'approved' and not pmdc_verified from doctor_profiles where id = :i", i=new_doctor_profile) is True)

    print("\n[17d] Email: sign-up verification, forgot password, appointment emails and reminders")
    from datetime import datetime as _dt, timedelta as _td, timezone as _tz

    from auth import hash_password as _hash_password
    from services import auth_tokens as _tokens
    from services import email_outbox as _outbox

    def token_rows(user_id: int, purpose: str) -> int:
        return scalar("select count(*) from auth_tokens where user_id = :u and purpose = :p", u=user_id, p=purpose)

    def put_code(user_id: int, code: str, *, minutes: int = 10, attempts: int = 0) -> None:
        """Store a reset code the test knows (the real one is only ever in the email)."""
        with engine.begin() as conn:
            conn.execute(text("update auth_tokens set used_at = now() where user_id = :u and purpose = 'reset_password' and used_at is null"), {"u": user_id})
            conn.execute(text("insert into auth_tokens (user_id, purpose, token_hash, attempts, expires_at, created_at) "
                              "values (:u, 'reset_password', :h, :a, :e, now() - interval '2 minutes')"),
                         {"u": user_id, "h": _tokens.fingerprint("reset_password", code, user_id), "a": attempts,
                          "e": _dt.now(_tz.utc) + _td(minutes=minutes)})

    # -- sign-up: unverified until the link is used ----------------------------------------------
    mail_email = f"{TEST_PREFIX}mail-{run_id}@example.com"
    response = call("POST", "/api/auth/signup", body={
        "role": "patient", "full_name": f"{TEST_PREFIX}mail user", "email": mail_email, "password": "Mail-test-pass-1",
        "confirm_password": "Mail-test-pass-1", "phone": f"0313 {secrets.randbelow(10**7):07d}", "cnic": "12345-1234567-1",
        "date_of_birth": "1995-05-05", "gender": "female"})
    check("sign-up works as before", response.status_code == 201, f"got {response.status_code}")
    client.cookies.clear()  # the sign-up route set the new user's session cookies on the shared client: later "no token" calls must be anonymous
    mail_user = scalar("select id from users where email = :e", e=mail_email)
    check("a new account starts with an unconfirmed email", scalar("select is_email_verified from users where id = :i", i=mail_user) is False)
    check("and has one verification link waiting", token_rows(mail_user, "verify_email") == 1)
    check("sign-up wrote no content to the outbox (nothing is sent without a provider)", scalar("select count(*) from email_outbox where user_id = :u", u=mail_user) == 0)

    check("confirming with a made-up token -> 400", call("POST", "/api/backend/auth/verify-email", None, {"token": "x" * 43}).status_code == 400)
    check("a too-short token -> 422", call("POST", "/api/backend/auth/verify-email", None, {"token": "abc"}).status_code == 422)
    link_token = "T" * 43
    with engine.begin() as conn:  # a link token the test knows
        conn.execute(text("update auth_tokens set used_at = now() where user_id = :u and purpose = 'verify_email'"), {"u": mail_user})
        conn.execute(text("insert into auth_tokens (user_id, purpose, token_hash, expires_at) values (:u, 'verify_email', :h, now() + interval '1 day')"),
                     {"u": mail_user, "h": _tokens.fingerprint("verify_email", link_token)})
    response = call("POST", "/api/backend/auth/verify-email", None, {"token": link_token})
    check("the link confirms the email", response.status_code == 200 and scalar("select is_email_verified from users where id = :i", i=mail_user) is True,
          f"got {response.status_code}")
    check("a link works once", call("POST", "/api/backend/auth/verify-email", None, {"token": link_token}).status_code == 400)
    with engine.begin() as conn:
        conn.execute(text("insert into auth_tokens (user_id, purpose, token_hash, expires_at) values (:u, 'verify_email', :h, now() - interval '1 minute')"),
                     {"u": mail_user, "h": _tokens.fingerprint("verify_email", "E" * 43)})
    check("an expired link -> 400", call("POST", "/api/backend/auth/verify-email", None, {"token": "E" * 43}).status_code == 400)

    no_token = call("POST", "/api/backend/auth/resend-verification")
    check("resend needs a sign-in -> 401", no_token.status_code == 401, f"got {no_token.status_code} {no_token.text[:80]}")
    check("resend for an already confirmed address sends nothing and says so",
          "already confirmed" in call("POST", "/api/backend/auth/resend-verification", tok_a).json().get("message", "")
          and token_rows(patient_a, "verify_email") == 0)
    with engine.begin() as conn:
        conn.execute(text("update users set is_email_verified = false where id = :i"), {"i": patient_a})
    first = call("POST", "/api/backend/auth/resend-verification", tok_a)
    check("an unconfirmed account can ask for a new link", first.status_code == 200 and token_rows(patient_a, "verify_email") == 1, f"got {first.status_code}")
    again = call("POST", "/api/backend/auth/resend-verification", tok_a)
    check("asking again at once -> 429 with how long to wait", again.status_code == 429 and again.headers.get("retry-after", "0").isdigit()
          and token_rows(patient_a, "verify_email") == 1, f"got {again.status_code}")
    with engine.begin() as conn:
        conn.execute(text("update users set is_email_verified = true where id = :i"), {"i": patient_a})
        conn.execute(text("delete from auth_tokens where user_id = :i"), {"i": patient_a})

    # -- forgot password -------------------------------------------------------------------------
    forgot = "/api/backend/auth/forgot-password"
    unknown = call("POST", forgot, None, {"email": f"{TEST_PREFIX}nobody-{run_id}@example.com"})
    known = call("POST", forgot, None, {"email": mail_email})
    check("forgot password answers the same for an unknown and a known address",
          unknown.status_code == 202 == known.status_code and unknown.json() == known.json(), f"got {unknown.status_code} {known.status_code}")
    check("only the real account gets a code", scalar("select count(*) from users u join auth_tokens t on t.user_id = u.id where u.email like :p and t.purpose = 'reset_password'",
                                                   p=f"{TEST_PREFIX}nobody-%") == 0 and token_rows(mail_user, "reset_password") == 1)
    again = call("POST", forgot, None, {"email": mail_email})
    check("asking again within a minute is answered the same and makes no second code", again.status_code == 202 and token_rows(mail_user, "reset_password") == 1)
    check("a malformed address -> 422", call("POST", forgot, None, {"email": "not-an-email"}).status_code == 422)
    inactive = create_user("mail-inactive", "PATIENT", active=False)
    with engine.begin() as conn:
        email_of_inactive = conn.execute(text("select email from users where id = :i"), {"i": inactive}).scalar()
    call("POST", forgot, None, {"email": email_of_inactive})
    check("an inactive account gets no code", token_rows(inactive, "reset_password") == 0)

    # -- reset password --------------------------------------------------------------------------
    reset = "/api/backend/auth/reset-password"
    new_password = "Brand-new-pass-9"
    body = {"email": mail_email, "code": "123456", "new_password": new_password, "confirm_password": new_password}
    put_code(mail_user, "123456")
    check("a code that is not 6 digits -> 422", call("POST", reset, None, {**body, "code": "12ab56"}).status_code == 422)
    check("passwords that differ -> 422", call("POST", reset, None, {**body, "confirm_password": "other-pass-1"}).status_code == 422)
    check("a short password -> 422", call("POST", reset, None, {**body, "new_password": "short", "confirm_password": "short"}).status_code == 422)
    wrong = call("POST", reset, None, {**body, "code": "654321"})
    check("a wrong code -> 400, the same answer as an unknown address",
          wrong.status_code == 400 and call("POST", reset, None, {**body, "email": f"{TEST_PREFIX}nobody-{run_id}@example.com"}).json() == wrong.json())
    check("each wrong try is counted", scalar("select attempts from auth_tokens where user_id = :u and purpose = 'reset_password' and used_at is null", u=mail_user) == 1)

    old_login = httpx.post(f"{BASE_URL}/api/auth/login", json={"email": mail_email, "password": "Mail-test-pass-1"}, timeout=60)
    old_refresh = scalar("select count(*) from refresh_tokens where user_id = :u and revoked_at is null", u=mail_user)
    check("before the reset the old password signs in and a session exists", old_login.status_code == 200 and old_refresh >= 1, f"got {old_login.status_code}")
    put_code(mail_user, "123456")
    done = call("POST", reset, None, body)
    check("the right code sets the new password", done.status_code == 200, f"got {done.status_code} {done.text[:80]}")
    check("every session of the account is signed out", scalar("select count(*) from refresh_tokens where user_id = :u and revoked_at is null", u=mail_user) == 0)
    check("the old password no longer works", httpx.post(f"{BASE_URL}/api/auth/login", json={"email": mail_email, "password": "Mail-test-pass-1"}, timeout=60).status_code == 401)
    check("the new password works", httpx.post(f"{BASE_URL}/api/auth/login", json={"email": mail_email, "password": new_password}, timeout=60).status_code == 200)
    check("a code works once", call("POST", reset, None, body).status_code == 400)

    put_code(mail_user, "111111", attempts=5)
    check("after 5 wrong tries even the right code is refused", call("POST", reset, None, {**body, "code": "111111"}).status_code == 400)
    put_code(mail_user, "222222", minutes=-1)
    check("an expired code is refused", call("POST", reset, None, {**body, "code": "222222"}).status_code == 400)
    with engine.begin() as conn:
        conn.execute(text("update users set is_email_verified = false where id = :i"), {"i": mail_user})
    put_code(mail_user, "333333")
    call("POST", reset, None, {**body, "code": "333333", "new_password": new_password + "x", "confirm_password": new_password + "x"})
    check("a reset also confirms the address (the code reached the inbox)", scalar("select is_email_verified from users where id = :i", i=mail_user) is True)
    check("the code is stored hashed, never in the clear",
          scalar("select count(*) from auth_tokens where token_hash in ('123456', '333333', '111111')") == 0
          and scalar("select min(length(token_hash)) from auth_tokens where user_id = :u", u=mail_user) == 64)

    # -- appointment emails and reminders --------------------------------------------------------
    def outbox(kind: str, appointment_id: int, status: str = "pending") -> int:
        return scalar("select count(*) from email_outbox where kind = :k and appointment_id = :a and status = :s", k=kind, a=appointment_id, s=status)

    when = _dt.fromisoformat(next_slot(doctor).replace("Z", "+00:00"))
    with engine.begin() as conn:
        conn.execute(text("update users set email_reminders = true where id in (:a, :d)"), {"a": patient_a, "d": doctor})
    response = call("POST", "/api/backend/appointments", tok_a, {"doctor_id": doctor, "date": when.isoformat(), "mode": "in_clinic", "reason": "private chest pain details"})
    booked_mail = response.json().get("appointment", {}) if response.status_code == 201 else {}
    appt = booked_mail.get("id", 0)
    check("booking for the email checks -> 201", response.status_code == 201, f"got {response.status_code} {response.text[:100]}")
    check("a confirmation is queued for the patient and for the doctor", outbox("appointment_booked", appt) == 2
          and scalar("select count(distinct user_id) from email_outbox where appointment_id = :a and kind = 'appointment_booked'", a=appt) == 2)
    check("a reminder is queued for each, five minutes before the start", outbox("appointment_reminder", appt) == 2
          and scalar("select count(*) from email_outbox where appointment_id = :a and kind = 'appointment_reminder' and send_after = :t", a=appt, t=when - _td(minutes=5)) == 2)
    check("each reminder expires three minutes after it is due",
          scalar("select count(*) from email_outbox where appointment_id = :a and kind = 'appointment_reminder' and expires_at = send_after + interval '3 minutes'", a=appt) == 2)
    check("the queued rows hold times and names but never the reason",
          scalar("select count(*) from email_outbox where appointment_id = :a and (context::text ilike '%chest%' or context::text ilike '%private%')", a=appt) == 0)
    check("the queue sent nothing: no provider is configured in tests", scalar("select count(*) from email_outbox where appointment_id = :a and status = 'sent'", a=appt) == 0)

    moved_to = _dt.fromisoformat(next_slot(doctor).replace("Z", "+00:00"))
    response = call("PUT", f"/api/backend/appointments/{appt}", tok_a, {"date": moved_to.isoformat()})
    check("moving the appointment -> 200", response.status_code == 200, f"got {response.status_code}")
    check("the old reminders are withdrawn and new ones queued for the new time",
          outbox("appointment_reminder", appt, "cancelled") == 2 and outbox("appointment_reminder", appt) == 2
          and scalar("select count(*) from email_outbox where appointment_id = :a and kind = 'appointment_reminder' and status = 'pending' and send_after = :t", a=appt, t=moved_to - _td(minutes=5)) == 2)
    check("both are told about the move", outbox("appointment_moved", appt) == 2)
    check("the move email carries the old time",
          scalar("select count(*) from email_outbox where appointment_id = :a and kind = 'appointment_moved' and context->>'old_when_text' is not null", a=appt) == 2)

    response = call("PUT", f"/api/backend/appointments/{appt}", tok_a, {"status": "CANCELLED"})
    check("cancelling -> 200", response.status_code == 200, f"got {response.status_code}")
    check("pending reminders are withdrawn and both are told", outbox("appointment_reminder", appt) == 0 and outbox("appointment_cancelled", appt) == 2)
    check("the cancel email says who cancelled",
          scalar("select count(*) from email_outbox where appointment_id = :a and kind = 'appointment_cancelled' and context->>'cancelled_by' = 'patient'", a=appt) == 2)

    # an unconfirmed address, a person who switched reminders off, and a booking made minutes before the start
    with engine.begin() as conn:
        conn.execute(text("update users set email_reminders = false where id = :i"), {"i": doctor_2})
        conn.execute(text("update users set is_email_verified = false where id = :i"), {"i": patient_b})
    later = _dt.fromisoformat(next_slot(doctor_2).replace("Z", "+00:00"))
    response = call("POST", "/api/backend/appointments", token_for(patient_b), {"doctor_id": doctor_2, "date": later.isoformat(), "mode": "in_clinic"})
    second = response.json().get("appointment", {}).get("id", 0) if response.status_code == 201 else 0
    check("booking with an unconfirmed patient and a doctor who turned reminders off -> 201", response.status_code == 201, f"got {response.status_code}")
    check("the unconfirmed patient gets no email at all; the doctor gets the confirmation only",
          scalar("select count(*) from email_outbox where appointment_id = :a and user_id = :p", a=second, p=patient_b) == 0
          and outbox("appointment_booked", second) == 1 and outbox("appointment_reminder", second) == 0)
    with engine.begin() as conn:
        conn.execute(text("update users set email_reminders = true where id = :i"), {"i": doctor_2})
        conn.execute(text("update users set is_email_verified = true where id = :i"), {"i": patient_b})

    # the reminders switch
    response = call("PATCH", "/api/backend/profile/account", tok_a, {"full_name": f"{TEST_PREFIX}patient a", "phone": scalar("select phone from users where id = :i", i=patient_a) or "0300 0000001", "email_reminders": False})
    check("a person can switch reminder emails off in their profile", response.status_code == 200
          and scalar("select email_reminders from users where id = :i", i=patient_a) is False
          and call("GET", "/api/backend/profile", tok_a).json()["account"]["email_reminders"] is False, f"got {response.status_code} {response.text[:100]}")
    with engine.begin() as conn:
        conn.execute(text("update users set email_reminders = true where id = :i"), {"i": patient_a})

    # -- the sender: due rows go once, late and stale reminders do not ---------------------------
    from database import SessionLocal as _Session

    sent_log: list[dict] = []

    def fake_sender(**kwargs):
        sent_log.append(kwargs)
        return f"em_{len(sent_log)}"

    soon = _dt.now(_tz.utc) + _td(hours=3)
    with engine.begin() as conn:  # earlier sections queued rows too: only this section's rows should be due now
        conn.execute(text("update email_outbox set status = 'cancelled' where status = 'pending'"))
        conn.execute(text("update appointments set status = 'BOOKED', date = :d where id = :i"), {"d": soon, "i": appt})
        for n, (kind, offset_s, expires_s, ctx_start) in enumerate([
                ("appointment_reminder", -30, 120, soon),                       # due now, in time, still true
                ("appointment_reminder", -400, -60, soon),                      # missed its window: expired
                ("appointment_reminder", -30, 120, soon + _td(days=1))]):       # the appointment has since moved: stale
            conn.execute(text("insert into email_outbox (user_id, to_address, kind, context, dedupe_key, appointment_id, status, send_after, expires_at) "
                              "values (:u, :to, :k, cast(:c as json), :key, :a, 'pending', now() + make_interval(secs => :o), now() + make_interval(secs => :x))"),
                         {"u": patient_a, "to": f"{TEST_PREFIX}a-{run_id}@example.test", "k": kind, "a": appt, "key": f"{TEST_PREFIX}dispatch-{run_id}-{n}",
                          "o": offset_s, "x": expires_s,
                          "c": json.dumps({"audience": "patient", "recipient_name": "A", "patient_name": "A", "doctor_name": "D", "when_text": "Mon 1 Jan, 9:00 AM",
                                           "mode": "in_clinic", "start_iso": ctx_start.isoformat(), "minutes": 5})})
    with _Session() as session:
        went = _outbox.dispatch_due(session, sender=fake_sender, pause=0)
    check("the sender sends the one reminder that is due and true", went == 1 and len(sent_log) == 1, f"sent {went}")
    check("it sends an email with both parts, from the queue's own key (so a retry cannot double it)",
          sent_log and sent_log[0]["html"] and sent_log[0]["text"] and sent_log[0]["idempotency_key"] == f"{TEST_PREFIX}dispatch-{run_id}-0")
    check("a reminder that missed its window is dropped, not sent late",
          scalar("select status from email_outbox where dedupe_key = :k", k=f"{TEST_PREFIX}dispatch-{run_id}-1") == "expired")
    check("a reminder for a time the appointment no longer has is cancelled",
          scalar("select status from email_outbox where dedupe_key = :k", k=f"{TEST_PREFIX}dispatch-{run_id}-2") == "cancelled")
    with _Session() as session:
        check("running it again sends nothing more", _outbox.dispatch_due(session, sender=fake_sender, pause=0) == 0 and len(sent_log) == 1)

    def failing(retryable: bool):
        def send_failing(**kwargs):
            from services.email_client import EmailError
            raise EmailError("HTTP 429 rate_limit_exceeded" if retryable else "HTTP 422 validation_error", retryable=retryable)
        return send_failing

    with engine.begin() as conn:
        for n, retryable in enumerate((True, False)):
            conn.execute(text("insert into email_outbox (user_id, to_address, kind, context, dedupe_key, status, send_after) "
                              "values (:u, :to, 'password_changed', cast(:c as json), :key, 'pending', now() - interval '1 second')"),
                         {"u": patient_a, "to": f"{TEST_PREFIX}b-{run_id}@example.test", "key": f"{TEST_PREFIX}fail-{run_id}-{n}", "c": json.dumps({"name": "A"})})
    with _Session() as session:
        _outbox.dispatch_due(session, sender=failing(True), pause=0, limit=50)
    rows = {r[0]: r[1:] for r in engine.connect().execute(text(
        "select dedupe_key, status, attempts, last_error from email_outbox where dedupe_key like :p"), {"p": f"{TEST_PREFIX}fail-{run_id}-%"}).fetchall()}
    check("a provider error that may pass is retried later, not dropped",
          all(v[0] == "pending" and v[1] == 1 and "429" in v[2] for v in rows.values()), f"got {rows}")
    check("and waits before the next try", scalar("select min(send_after) > now() from email_outbox where dedupe_key like :p", p=f"{TEST_PREFIX}fail-{run_id}-%") is True)
    with engine.begin() as conn:
        conn.execute(text("update email_outbox set send_after = now() - interval '1 second' where dedupe_key like :p"), {"p": f"{TEST_PREFIX}fail-{run_id}-%"})
    with _Session() as session:
        _outbox.dispatch_due(session, sender=failing(False), pause=0, limit=50)
    check("an error that will never pass marks the email failed at once",
          scalar("select count(*) from email_outbox where dedupe_key like :p and status = 'failed'", p=f"{TEST_PREFIX}fail-{run_id}-%") == 2)
    from services.email_client import EmailDisabled as _Disabled

    def no_provider(**kwargs):
        raise _Disabled

    with engine.begin() as conn:
        conn.execute(text("insert into email_outbox (user_id, to_address, kind, context, dedupe_key, status, send_after) "
                          "values (:u, :to, 'password_changed', cast(:c as json), :key, 'pending', now() - interval '1 second')"),
                     {"u": patient_a, "to": f"{TEST_PREFIX}c-{run_id}@example.test", "key": f"{TEST_PREFIX}nop-{run_id}", "c": json.dumps({"name": "A"})})
    with _Session() as session:
        _outbox.dispatch_due(session, sender=no_provider, pause=0)
    check("without a provider the queue waits and counts no attempt",
          scalar("select status = 'pending' and attempts = 0 from email_outbox where dedupe_key = :k", k=f"{TEST_PREFIX}nop-{run_id}") is True)

    with engine.begin() as conn:
        conn.execute(text("update appointments set status = 'CANCELLED' where id = :i"), {"i": appt})

    print("\n[17e] In-app notifications: events, links, filters, unread count, dismiss")

    def told(user_id: int, kind: str) -> int:
        return scalar("select count(*) from notifications where user_id = :u and kind = :k", u=user_id, k=kind)

    note_path = "/api/backend/notifications"
    check("the unread count needs a token -> 401", call("GET", note_path + "/unread-count").status_code == 401)
    count_now = call("GET", note_path + "/unread-count", tok_c)
    check("the unread count is one small number", count_now.status_code == 200 and set(count_now.json()) == {"unread"}
          and count_now.json()["unread"] == call("GET", note_path, tok_c).json()["unread"], f"got {count_now.status_code}")

    # events that happened earlier in this run, told to the right person with a page to go to
    check("a doctor opening a patient's file is told to the patient (once per unlock, with a link to the access history)",
          told(patient_a, "access_opened") >= 1
          and scalar("select count(*) from notifications where user_id = :u and kind = 'access_opened' and link = '/dashboard/patient#access-history'", u=patient_a) >= 1)
    check("a doctor's change to the health record is told to the patient", told(patient_c, "record_changed") >= 2
          and scalar("select count(*) from notifications where user_id = :u and kind = 'record_changed' and link = '/dashboard/patient#health-record'", u=patient_c) >= 2)
    check("a new prescription and its cancellation are told to the patient", told(patient_c, "prescription_added") == 1 and told(patient_c, "prescription_cancelled") == 1)
    check("a notification never holds a medicine name or a diagnosis",
          scalar("select count(*) from notifications where user_id = :u and kind like 'prescription%' and (body ilike '%mg%' or body ilike '%warfarin%' or body ilike '%metformin%')", u=patient_c) == 0
          and scalar("select count(*) from notifications where user_id = :u and kind = 'record_changed' and (body ilike '%penicillin%' or body ilike '%diabetes%')", u=patient_c) == 0)
    check("a new doctor signing up is told to the administrators, with a link to the review queue",
          scalar("select count(*) from notifications where user_id = :a and kind = 'admin_doctor_review' and link = '/admin#doctors'", a=admin) >= 1)
    check("approving a doctor tells the doctor", told(scalar("select user_id from doctor_profiles where id = :i", i=new_doctor_profile), "doctor_approved") == 1)
    check("giving and taking away the PMDC badge each tell the doctor", told(doctor, "doctor_pmdc") == 2)
    check("each password reset tells the account (two succeeded above)", told(mail_user, "account_password") == 2)

    # filters, links and the fallback link for an older notification
    listed = call("GET", note_path + "?limit=100", tok_c).json()
    check("every notification says its category and where it leads",
          all(n["category"] in {"appointments", "payments", "records", "account", "admin"} for n in listed["items"])
          and any(n["link"] == "/dashboard/patient#prescriptions" for n in listed["items"]))
    records_only = call("GET", note_path + "?category=records&limit=100", tok_c).json()
    check("the category filter returns only that category", records_only["total"] >= 3 and all(n["category"] == "records" for n in records_only["items"]))
    check("an unknown category -> 422", call("GET", note_path + "?category=gossip", tok_c).status_code == 422)
    check("the unread count counts everything, not only the filter", records_only["unread"] == listed["unread"])
    with engine.begin() as conn:  # an older notification has no stored link
        legacy_id = conn.execute(text("insert into notifications (user_id, title, body, status, kind) values (:u, 'Appointment confirmed', 'older one', 'UNREAD', 'appointment_confirmed') returning id"), {"u": patient_c}).scalar_one()
    legacy = next(n for n in call("GET", note_path + "?limit=100", tok_c).json()["items"] if n["id"] == legacy_id)
    check("an older notification still leads somewhere (by its kind and the reader's role)", legacy["link"] == "/doctors")
    call("POST", f"{note_path}/read-all", tok_c)
    check("read-all leaves nothing unread and the unread filter is then empty",
          call("GET", note_path + "/unread-count", tok_c).json()["unread"] == 0 and call("GET", note_path + "?unread_only=true", tok_c).json()["total"] == 0)
    check("reading sets the time it was read", scalar("select count(*) from notifications where user_id = :u and status = 'READ' and read_at is null", u=patient_c) == 0)

    # dismissing: your own only
    check("another user cannot dismiss my notification -> 404", call("DELETE", f"{note_path}/{legacy_id}", tok_a).status_code == 404)
    check("dismissing needs a token -> 401", call("DELETE", f"{note_path}/{legacy_id}").status_code == 401)
    gone = call("DELETE", f"{note_path}/{legacy_id}", tok_c)
    check("I can dismiss my own notification", gone.status_code == 200 and scalar("select count(*) from notifications where id = :i", i=legacy_id) == 0, f"got {gone.status_code}")
    check("dismissing it twice -> 404", call("DELETE", f"{note_path}/{legacy_id}", tok_c).status_code == 404)

    # an event delivered twice is told once
    from services import notifications as _notes
    from database import SessionLocal as _NoteSession

    with _NoteSession() as session:
        first = _notes.add(session, patient_c, "payment_received", "Payment received", "x", dedupe_key=f"{TEST_PREFIX}dedupe-{run_id}")
        session.commit()
        again = _notes.add(session, patient_c, "payment_received", "Payment received", "x", dedupe_key=f"{TEST_PREFIX}dedupe-{run_id}")
        session.commit()
    check("an event with the same key is told once", first is not None and again is None
          and scalar("select count(*) from notifications where dedupe_key = :k", k=f"{TEST_PREFIX}dedupe-{run_id}") == 1)
    with engine.begin() as conn:
        conn.execute(text("delete from notifications where dedupe_key = :k"), {"k": f"{TEST_PREFIX}dedupe-{run_id}"})

    print("\n[18] Sessions: HttpOnly cookies, refresh-token rotation, replay detection, logout")
    from auth import hash_password

    password = "Session-test-pass-1"
    session_user = create_user("session", "PATIENT")
    with engine.begin() as conn:
        conn.execute(text("update users set hashed_password = :pw, email = :email where id = :id"),
                     {"pw": hash_password(password), "email": f"{TEST_PREFIX}session-{run_id}@example.com", "id": session_user})
    session_email = f"{TEST_PREFIX}session-{run_id}@example.com"

    response = call("POST", "/api/auth/login", body={"email": session_email, "password": password})
    set_cookies = response.headers.get_list("set-cookie")
    check("login through the website -> 200 with only the user in the body",
          response.status_code == 200 and set(response.json()) == {"user"} and response.json()["user"]["email"] == session_email,
          f"got {response.status_code} {list(response.json())}")
    cookie_names = {c.split("=", 1)[0] for c in set_cookies}
    check("the access and refresh tokens are set as cookies", {"nora_at", "nora_rt"} <= cookie_names, f"got {cookie_names}")
    check("both cookies are HttpOnly and SameSite (JavaScript cannot read them)",
          len(set_cookies) == 2 and all("HttpOnly" in c and "SameSite=Lax" in c for c in set_cookies), f"got {set_cookies}")
    check("neither token appears in the response body", "eyJ" not in response.text and "token" not in response.text)

    cookies = {c.split("=", 1)[0]: c.split("=", 1)[1].split(";", 1)[0] for c in set_cookies}
    me = httpx.get(f"{BASE_URL}/api/backend/api/auth/me", headers={"Cookie": f"nora_at={cookies['nora_at']}"})
    check("the access cookie authenticates a request through the proxy -> 200", me.status_code == 200, f"got {me.status_code}")
    refreshed = httpx.get(f"{BASE_URL}/api/backend/api/auth/me", headers={"Cookie": f"nora_at=expired; nora_rt={cookies['nora_rt']}"})
    new_cookies = refreshed.headers.get_list("set-cookie")
    check("a bad access cookie with a valid refresh cookie is renewed in place -> 200 and new cookies",
          refreshed.status_code == 200 and len(new_cookies) == 2, f"got {refreshed.status_code} {len(new_cookies)}")
    time.sleep(4)  # past the 3-second grace for requests that raced the rotation
    replay = httpx.get(f"{BASE_URL}/api/backend/api/auth/me", headers={"Cookie": f"nora_at=expired; nora_rt={cookies['nora_rt']}"})
    check("the used refresh token presented again -> 401 (and its cookies are cleared)",
          replay.status_code == 401 and all("Max-Age=0" in c for c in replay.headers.get_list("set-cookie")), f"got {replay.status_code}")
    newest = {c.split("=", 1)[0]: c.split("=", 1)[1].split(";", 1)[0] for c in new_cookies}
    after_replay = httpx.get(f"{BASE_URL}/api/backend/api/auth/me", headers={"Cookie": f"nora_at=expired; nora_rt={newest['nora_rt']}"})
    check("the replay revoked the whole session: even the newest refresh token is dead -> 401", after_replay.status_code == 401,
          f"got {after_replay.status_code}")

    response = call("POST", "/api/backend/auth/login", body={"email": session_email, "password": password})
    check("the proxy does not expose login (tokens only leave through the cookie routes) -> 404", response.status_code == 404,
          f"got {response.status_code}")
    check("a cross-site browser post through the proxy is refused -> 403",
          httpx.post(f"{BASE_URL}/api/backend/appointments", headers={"Origin": "https://evil.example"}, json={}).status_code == 403)

    login = call("POST", "/api/auth/login", body={"email": session_email, "password": password})
    login_cookies = {c.split("=", 1)[0]: c.split("=", 1)[1].split(";", 1)[0] for c in login.headers.get_list("set-cookie")}
    out = httpx.post(f"{BASE_URL}/api/auth/logout", headers={"Cookie": f"nora_rt={login_cookies['nora_rt']}"})
    check("logout clears both cookies -> 200", out.status_code == 200 and all("Max-Age=0" in c for c in out.headers.get_list("set-cookie")),
          f"got {out.status_code}")
    gone = httpx.get(f"{BASE_URL}/api/backend/api/auth/me", headers={"Cookie": f"nora_at=expired; nora_rt={login_cookies['nora_rt']}"})
    check("after logout the refresh token no longer works -> 401", gone.status_code == 401, f"got {gone.status_code}")

    direct = httpx.post(f"{BACKEND_URL}/auth/login", json={"email": session_email, "password": password})
    first_refresh = direct.json().get("refresh_token", "")
    check("FastAPI login returns both tokens (for the Next.js server, never for browsers)",
          direct.status_code == 200 and first_refresh and direct.json().get("access_token"), f"got {direct.status_code}")
    with engine.begin() as conn:
        conn.execute(text("update users set is_active = false where id = :id"), {"id": session_user})
    inactive = httpx.post(f"{BACKEND_URL}/auth/refresh", json={"refresh_token": first_refresh})
    check("an account turned inactive cannot refresh -> 401", inactive.status_code == 401, f"got {inactive.status_code}")
    check("a made-up refresh token -> 401", httpx.post(f"{BACKEND_URL}/auth/refresh", json={"refresh_token": "x" * 40}).status_code == 401)
    with engine.begin() as conn:
        conn.execute(text("delete from refresh_tokens where user_id = :id"), {"id": session_user})

    client.cookies.clear()  # section 18 signed this client in through the website; later checks use tokens only
    print("\n[19] Profiles (Module 1): own profile only, verified fields locked, public doctor card")
    response = call("GET", "/api/backend/profile", tok_a)
    mine = response.json() if response.status_code == 200 else {}
    check("a patient reads their own profile (age, masked CNIC, no doctor part)",
          response.status_code == 200 and mine["account"]["id"] == patient_a and mine["patient"]["age"] >= 0
          and mine["patient"]["cnic_masked"].startswith("*****-") and mine["doctor"] is None, f"got {response.status_code}")
    check("profile without a token -> 401", call("GET", "/api/backend/profile").status_code == 401)
    form = {"date_of_birth": "1990-05-17", "gender": "female", "city": "Lahore", "blood_group": "o+",
            "emergency_contact_name": "Test Contact", "emergency_contact_phone": "0300 7654321", "user_id": patient_b}
    response = call("PUT", "/api/backend/patient/profile", tok_a, form)
    saved = response.json() if response.status_code == 200 else {}
    check("a patient saves details (body user_id=B ignored, gender and blood group normalised)",
          response.status_code == 200 and saved["patient"]["gender"] == "Female" and saved["patient"]["blood_group"] == "O+"
          and saved["account"]["id"] == patient_a, f"got {response.status_code}")
    check("B's profile is untouched", call("GET", "/api/backend/profile", tok_b).json()["patient"]["city"] is None)
    check("a future birth date -> 422",
          call("PUT", "/api/backend/patient/profile", tok_a, {**form, "date_of_birth": "2999-01-01"}).status_code == 422)
    check("a doctor saving a patient profile -> 403", call("PUT", "/api/backend/patient/profile", tok_doctor, form).status_code == 403)
    response = call("PATCH", "/api/backend/profile/account", tok_a, {"full_name": f"{TEST_PREFIX}patient A renamed", "phone": "0322 1112223"})
    check("a user changes their own name and phone",
          response.status_code == 200 and response.json()["account"]["phone"] == "0322 1112223", f"got {response.status_code}")
    check("a phone number already used by another account -> 409",
          call("PATCH", "/api/backend/profile/account", tok_b, {"full_name": f"{TEST_PREFIX}patient B", "phone": "0322 1112223"}).status_code == 409)
    pmdc_before = scalar("select pmdc_number from doctor_profiles where user_id = :id", id=doctor)
    response = call("PUT", "/api/backend/doctor/profile", tok_doctor,
                    {**doctor_form, "specialization": "Cardiology", "pmdc_number": "HACKED", "verification_status": "approved"})
    check("a doctor cannot change specialization or the PMDC number",
          response.status_code == 200 and response.json()["doctor"]["specialization"] == "General Physician"
          and scalar("select pmdc_number from doctor_profiles where user_id = :id", id=doctor) == pmdc_before, f"got {response.status_code}")
    card = next((d for d in call("GET", "/api/backend/appointments/doctors").json() if d["doctor_id"] == doctor), {})
    check("the public doctor card shows qualifications, experience, languages, modes and fee",
          card.get("qualifications") == "MBBS, FCPS" and card.get("experience_years") == 8
          and card.get("languages") == ["Urdu", "English"] and card.get("consultation_modes") == ["online", "in_clinic"]
          and card.get("consultation_fee") == 1500, f"got {card}")
    check("the public card never carries the PMDC number or phone", "pmdc_number" not in card and "phone" not in card)
    mine = call("GET", "/api/backend/profile", tok_a).json()["activity"]
    check("a patient's profile counts their own appointments and payments only",
          mine["appointments"] == scalar("select count(*) from appointments where user_id = :id", id=patient_a)
          and mine["payments"] == scalar("select count(*) from payments where user_id = :id", id=patient_a), f"got {mine}")
    theirs = call("GET", "/api/backend/profile", tok_doctor).json()["activity"]
    check("a doctor's profile counts appointments booked with them",
          theirs["appointments"] == scalar("select count(*) from appointments where doctor_id = :id", id=doctor), f"got {theirs}")

    client.cookies.clear()
    print("\n[20] Scheduling (Module 5): availability, slots, no double booking, reviews, notifications")
    sched_doctor = create_user("sched-doctor", "DOCTOR")
    tok_sched = token_for(sched_doctor)
    patient_s1 = create_user("sched-patient-1", "PATIENT")
    patient_s2 = create_user("sched-patient-2", "PATIENT")
    tok_s1, tok_s2 = token_for(patient_s1), token_for(patient_s2)

    response = call("GET", "/api/backend/doctor/availability", tok_sched)
    check("a doctor reads their weekly schedule (the test doctor has 7 windows, clinic time)",
          response.status_code == 200 and len(response.json()["windows"]) == 7 and response.json()["timezone"] == "Asia/Karachi",
          f"got {response.status_code}")
    check("a patient cannot read or set a doctor's availability",
          call("GET", "/api/backend/doctor/availability", tok_s1).status_code == 403
          and call("PUT", "/api/backend/doctor/availability", tok_s1, {"windows": []}).status_code == 403)
    window = {"weekday": 0, "start_time": "09:00:00", "end_time": "11:00:00", "slot_minutes": 30}
    bad = {"windows": [window, {**window, "start_time": "10:00:00", "end_time": "12:00:00"}]}
    check("overlapping windows -> 422 and the old schedule stays",
          call("PUT", "/api/backend/doctor/availability", tok_sched, bad).status_code == 422
          and len(call("GET", "/api/backend/doctor/availability", tok_sched).json()["windows"]) == 7)
    check("a 25-minute appointment length is refused -> 422",
          call("PUT", "/api/backend/doctor/availability", tok_sched, {"windows": [{**window, "slot_minutes": 25}]}).status_code == 422)
    all_week = {"windows": [{"weekday": d, "start_time": "09:00:00", "end_time": "11:00:00", "slot_minutes": 30} for d in range(7)]}
    response = call("PUT", "/api/backend/doctor/availability", tok_sched, all_week)
    check("the doctor publishes 09:00-11:00 every day in 30-minute slots", response.status_code == 200 and len(response.json()["windows"]) == 7)

    slots_url = f"/api/backend/appointments/doctors/{sched_doctor}/slots?days=3"
    open_slots = [s_ for d in call("GET", slots_url).json()["days"] for s_ in d["slots"] if s_["available"]]
    check("slots are public, 4 a day, and none is in the past or too soon",
          len(open_slots) >= 4 and all(datetime.fromisoformat(s_["start"].replace("Z", "+00:00")) > datetime.now(timezone.utc)
                                       for s_ in open_slots), f"got {len(open_slots)}")
    check("slots of an unverified doctor are not exposed -> 404",
          call("GET", f"/api/backend/appointments/doctors/{doctor_pending}/slots").status_code == 404)
    first, second, third = open_slots[0]["start"], open_slots[1]["start"], open_slots[2]["start"]

    response = call("POST", "/api/backend/appointments", tok_s1, {"doctor_id": sched_doctor, "date": first, "mode": "online"})
    mine = (response.json() or {}).get("appointment", {}) if response.status_code == 201 else {}
    check("a patient books an open slot -> 201 with its length and type",
          response.status_code == 201 and mine.get("duration_minutes") == 30 and mine.get("mode") == "online", f"got {response.status_code}")
    response = call("POST", "/api/backend/appointments", tok_s2, {"doctor_id": sched_doctor, "date": first})
    check("a second patient booking the same slot -> 409", response.status_code == 409, f"got {response.status_code} {response.text[:100]}")
    after = [s_ for d in call("GET", slots_url).json()["days"] for s_ in d["slots"] if s_["start"] == first]
    check("the taken slot is shown as not available", after and after[0]["available"] is False)
    odd = (datetime.fromisoformat(second.replace("Z", "+00:00")) + timedelta(minutes=7)).isoformat()
    check("a time that is not a slot -> 409", call("POST", "/api/backend/appointments", tok_s2,
                                                   {"doctor_id": sched_doctor, "date": odd}).status_code == 409)
    check("the doctor's only offered types are respected: an unlisted type -> 422",
          call("POST", "/api/backend/appointments", tok_s2, {"doctor_id": sched_doctor, "date": second, "mode": "teleport"}).status_code == 422)

    appt_url = f"/api/backend/appointments/{mine.get('id')}"
    response = call("PUT", appt_url, tok_s1, {"date": third})
    check("the patient moves their booking to another open slot",
          response.status_code == 200 and datetime.fromisoformat(response.json()["date"]) == datetime.fromisoformat(third.replace("Z", "+00:00")),
          f"got {response.status_code}")
    check("the old slot is open again", [s_["available"] for d in call("GET", slots_url).json()["days"]
                                         for s_ in d["slots"] if s_["start"] == first] == [True])
    check("another patient cannot move it -> 403", call("PUT", appt_url, tok_s2, {"date": second}).status_code == 403)
    check("the doctor cannot move a patient's booking -> 403", call("PUT", appt_url, tok_sched, {"date": second}).status_code == 403)
    check("an administrator does not change appointments -> 403", call("PUT", appt_url, tok_admin, {"status": "CANCELLED"}).status_code == 403)
    check("a visit that has not started cannot be completed -> 409", call("PUT", appt_url, tok_sched, {"status": "COMPLETED"}).status_code == 409)
    check("a patient cannot rate before it is completed -> 409",
          call("POST", appt_url + "/review", tok_s1, {"rating": 5}).status_code == 409)

    # The visit "happens": move it into the past directly in the database (we cannot wait for real time).
    with engine.begin() as conn:
        conn.execute(text("update appointments set date = now() - interval '2 hours' where id = :id"), {"id": mine.get("id")})
    response = call("PUT", appt_url, tok_sched, {"status": "COMPLETED"})
    check("the treating doctor completes it once it has started", response.status_code == 200 and response.json()["status"] == "COMPLETED",
          f"got {response.status_code}")
    check("a completed visit cannot be cancelled -> 409", call("PUT", appt_url, tok_s1, {"status": "CANCELLED"}).status_code == 409)
    check("another patient cannot rate it -> 403", call("POST", appt_url + "/review", tok_s2, {"rating": 1}).status_code == 403)
    check("a doctor cannot rate their own visit -> 403", call("POST", appt_url + "/review", tok_sched, {"rating": 5}).status_code == 403)
    check("a rating outside 1 to 5 -> 422", call("POST", appt_url + "/review", tok_s1, {"rating": 9}).status_code == 422)
    response = call("POST", appt_url + "/review", tok_s1, {"rating": 4, "comment": "  Clear and kind  "})
    check("the patient rates the completed visit once", response.status_code == 201 and response.json()["comment"] == "Clear and kind")
    check("rating twice -> 409", call("POST", appt_url + "/review", tok_s1, {"rating": 5}).status_code == 409)
    check("completing a visit asks the patient for a rating, once, with a link to their appointments",
          scalar("select count(*) from notifications where user_id = :u and kind = 'appointment_completed' and link = '/doctors'", u=patient_s1) == 1)
    check("a rating tells the doctor without saying who rated",
          scalar("select count(*) from notifications where user_id = :d and kind = 'review_received' and body not ilike '%' || :n || '%'", d=sched_doctor, n=f"{TEST_PREFIX}") == 1)
    card = next((d for d in call("GET", "/api/backend/appointments/doctors").json() if d["doctor_id"] == sched_doctor), {})
    check("the directory now shows the real average and count", card.get("rating") == 4.0 and card.get("rating_count") == 1, f"got {card.get('rating')}")
    reviews = call("GET", f"/api/backend/appointments/doctors/{sched_doctor}/reviews").json()
    check("reviews are public but carry no patient name", len(reviews) == 1 and set(reviews[0]) == {"rating", "comment", "created_at"})
    check("the patient's list says this visit was rated",
          next(a for a in call("GET", "/api/backend/appointments", tok_s1).json()["items"] if a["id"] == mine.get("id"))["reviewed"] is True)

    cancel_slot = open_slots[3]["start"]
    cancelled = call("POST", "/api/backend/appointments", tok_s2, {"doctor_id": sched_doctor, "date": cancel_slot}).json()["appointment"]
    check("the doctor can cancel a patient's booking", call("PUT", f"/api/backend/appointments/{cancelled['id']}", tok_sched, {"status": "CANCELLED"}).json()["status"] == "CANCELLED")
    check("a cancelled slot can be booked again",
          call("POST", "/api/backend/appointments", tok_s1, {"doctor_id": sched_doctor, "date": cancel_slot}).status_code == 201)

    response = call("GET", "/api/backend/notifications", tok_s1)
    page = response.json() if response.status_code == 200 else {}
    kinds = {n["kind"] for n in page.get("items", [])}
    check("the patient has a confirmation and a moved-appointment notice",
          {"appointment_confirmed", "appointment_rescheduled"} <= kinds and page.get("unread", 0) >= 2, f"got {kinds}")
    check("the doctor was told too", {"appointment_confirmed", "appointment_rescheduled"} <= {n["kind"] for n in call("GET", "/api/backend/notifications", tok_sched).json()["items"]})
    check("reminders for later are not shown yet",
          scalar("select count(*) from notifications where user_id = :id and kind = 'appointment_reminder' and scheduled_for > now()", id=patient_s1) > 0
          and "appointment_reminder" not in kinds)
    note = page["items"][0]
    check("another user cannot mark my notification read (looks like 404)",
          call("POST", f"/api/backend/notifications/{note['id']}/read", tok_s2).status_code == 404)
    check("I can mark one read", call("POST", f"/api/backend/notifications/{note['id']}/read", tok_s1).status_code == 200)
    call("POST", "/api/backend/notifications/read-all", tok_s1)
    check("read-all clears the unread count", call("GET", "/api/backend/notifications", tok_s1).json()["unread"] == 0)
    check("notifications need a token", call("GET", "/api/backend/notifications").status_code == 401)


def main() -> int:
    assert_local_databases()
    try:
        client.get("/", timeout=180)
    except httpx.HTTPError as error:
        sys.exit(f"ABORT: Next.js server not reachable at {BASE_URL} ({type(error).__name__}). Start it with `npm run dev`.")

    cleanup_test_users()  # leftovers from an interrupted earlier run
    try:
        run_tests()
    except Exception as error:  # report, then still clean up
        check("test run completed without errors", False, f"{type(error).__name__}: {error}")
    finally:
        cleanup_test_users()
        client.close()

    failed = [name for ok, name, _ in results if not ok]
    print(f"\n{len(results) - len(failed)} passed, {len(failed)} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
