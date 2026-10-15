"""Local OCR (services/local_ocr.py) and the choice of OCR engine (services/report_ocr.py).

The OCR engine is replaced with a fake, so these tests need no model download.
"""

from dataclasses import replace
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from services import local_ocr, report_ocr
from services.report_text import OcrError


class FakeOutput:
    """The shape RapidOCR 3 returns: boxes (corner points), texts and scores."""

    def __init__(self, cells):
        self.boxes = np.array([box for box, _, _ in cells], dtype=float) if cells else None
        self.txts = tuple(text for _, text, _ in cells)
        self.scores = tuple(score for _, _, score in cells)


def cell(text, score, x, y, width=None, height=20):
    """A text box with its top-left corner at (x, y)."""
    width = width or 12 * len(text)
    return [[x, y], [x + width, y], [x + width, y + height], [x, y + height]], text, score


class FakeEngine:
    def __init__(self, cells):
        self.cells = cells
        self.calls = 0

    def __call__(self, image):
        self.calls += 1
        return FakeOutput(self.cells)


@pytest.fixture
def fake_engine(monkeypatch: pytest.MonkeyPatch):
    def install(cells):
        engine = FakeEngine(cells)
        monkeypatch.setattr(local_ocr, "_engine", lambda: engine)
        return engine

    return install


def png(tmp_path: Path) -> Path:
    path = tmp_path / "scan.png"
    Image.new("RGB", (40, 40), "white").save(path)
    return path


# A one-page PDF with a text layer ("Hemoglobin 10.8 g/dL ..."), written by hand.
TEXT_PDF = b"""%PDF-1.4
1 0 obj <</Type/Catalog/Pages 2 0 R>> endobj
2 0 obj <</Type/Pages/Kids[3 0 R]/Count 1>> endobj
3 0 obj <</Type/Page/Parent 2 0 R/MediaBox[0 0 400 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>> endobj
4 0 obj <</Length 88>> stream
BT /F1 12 Tf 20 150 Td (Hemoglobin 10.8 g/dL. Platelets 230. Glucose 95 mg/dL.) Tj ET
endstream endobj
5 0 obj <</Type/Font/Subtype/Type1/BaseFont/Helvetica>> endobj
trailer <</Root 1 0 R>>
%%EOF
"""


def test_boxes_are_rebuilt_into_rows_left_to_right(tmp_path, fake_engine):
    # OCR returns the value cells before their names; rows put each value next to its test.
    fake_engine(
        [
            cell("10.8", 0.97, 300, 100),
            cell("g/dL", 0.95, 400, 101),
            cell("Hemoglobin", 0.99, 20, 99),
            cell("230", 0.96, 300, 140),
            cell("Platelets", 0.98, 20, 141),
        ]
    )
    read = local_ocr.read_text(png(tmp_path))
    assert read.text == "Hemoglobin 10.8 g/dL\nPlatelets 230"
    assert read.rows[0]["cells"][1] == {"text": "10.8", "score": 0.97}


def test_a_tilted_page_keeps_its_rows(tmp_path, fake_engine):
    # 3 degrees of tilt: the value of row 1 sits lower than the name of row 2 would if read flat.
    tilt = 0.0524
    cells = []
    for row, (name, value) in enumerate([("Hemoglobin", "10.8"), ("Platelets", "230")]):
        for x, text in [(20, name), (520, value)]:
            y = 100 + row * 22 + x * tilt
            width = 12 * len(text)
            box = [[x, y], [x + width, y + width * tilt], [x + width, y + width * tilt + 20], [x, y + 20]]
            cells.append((box, text, 0.97))
    fake_engine(cells)
    assert local_ocr.read_text(png(tmp_path)).text == "Hemoglobin 10.8\nPlatelets 230"


def test_logo_text_is_cleaned_and_low_score_words_dropped_but_numbers_never(tmp_path, fake_engine):
    fake_engine(
        [
            cell("IDC 鑫品 IMAGING", 0.9, 20, 20),
            cell("smudge", 0.2, 300, 60),
            cell("Hemoglobin", 0.99, 20, 100),
            cell("9.5", 0.41, 300, 100),  # unsure, but dropping it would let a reader take the next number
            cell("12-17.5", 0.98, 400, 100),
        ]
    )
    read = local_ocr.read_text(png(tmp_path))
    assert read.text == "IDC IMAGING\nHemoglobin 9.5 12-17.5"  # the number is kept exactly as read
    assert read.confidence == pytest.approx((0.9 + 0.99 + 0.41 + 0.98) / 4)


def test_pdf_text_layer_is_used_without_ocr(tmp_path, fake_engine):
    engine = fake_engine([("should not be used", 0.99)])
    path = tmp_path / "digital.pdf"
    path.write_bytes(TEXT_PDF)
    read = local_ocr.read_text(path)
    assert "Hemoglobin 10.8 g/dL" in read.text
    assert engine.calls == 0 and read.confidence is None and read.rows == []


def test_unreadable_file_is_a_local_ocr_error(tmp_path, fake_engine):
    fake_engine([])
    path = tmp_path / "broken.png"
    path.write_bytes(b"not an image")
    with pytest.raises(local_ocr.LocalOcrError):
        local_ocr.read_text(path)


def test_a_model_with_the_wrong_checksum_is_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(local_ocr, "settings", replace(local_ocr.settings, ocr_model_dir=str(tmp_path)))

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

        def raise_for_status(self):
            pass

        def iter_bytes(self):
            yield b"tampered model"

    monkeypatch.setattr(local_ocr.httpx, "stream", lambda *args, **kwargs: FakeResponse())
    with pytest.raises(local_ocr.LocalOcrError, match="checksum"):
        local_ocr.ensure_models()
    assert not any(tmp_path.iterdir())  # nothing left behind


@pytest.mark.anyio
async def test_uploads_are_read_with_local_ocr(tmp_path, fake_engine):
    fake_engine([cell("Hemoglobin", 0.97, 20, 100), cell("10.8", 0.97, 300, 100)])
    result = await report_ocr.extract_report_text(str(png(tmp_path)), "scan.png", None)
    assert result["structured_data_json"]["source"] == "local_ocr"
    assert "Hemoglobin 10.8" in result["cleaned_text"] and result["confidence"] == pytest.approx(0.97)
    assert result["structured_data_json"]["ocr_rows"][0]["text"] == "Hemoglobin 10.8"


@pytest.mark.anyio
async def test_no_text_found_is_an_error_not_an_empty_report(tmp_path, fake_engine):
    fake_engine([])
    with pytest.raises(OcrError) as caught:
        await report_ocr.extract_report_text(str(png(tmp_path)), "scan.png", None)
    assert caught.value.status_code == 422


# --- Word and TIFF ------------------------------------------------------------------------------


def make_docx(tmp_path: Path, picture: bool = False) -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph("Lab report")
    table = document.add_table(rows=2, cols=3)
    for row, values in enumerate([("Hemoglobin", "10.8", "g/dL"), ("Creatinine", "1.1", "mg/dL")]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    document.add_paragraph("Allergies: none known")
    if picture:
        image_path = tmp_path / "pasted.png"
        Image.new("RGB", (60, 60), "white").save(image_path)
        document.add_picture(str(image_path))
    path = tmp_path / "report.docx"
    document.save(path)
    return path


def test_word_text_and_tables_are_read_exactly_in_order(tmp_path):
    read = local_ocr.read_text(make_docx(tmp_path))
    assert read.text.splitlines() == [
        "Lab report",
        "Hemoglobin 10.8 g/dL",
        "Creatinine 1.1 mg/dL",
        "Allergies: none known",
    ]
    assert read.confidence is None and read.rows == []  # nothing was OCR'd: the text is exact


def test_a_picture_pasted_into_word_is_read_with_ocr(tmp_path, fake_engine):
    fake_engine([cell("Glucose", 0.95, 20, 100), cell("186", 0.95, 300, 100)])
    read = local_ocr.read_text(make_docx(tmp_path, picture=True))
    assert "Glucose 186" in read.text and read.text.startswith("Lab report")
    assert read.confidence == pytest.approx(0.95)


def test_every_page_of_a_tiff_is_read(tmp_path, fake_engine):
    engine = fake_engine([cell("Hemoglobin", 0.9, 20, 100), cell("10.8", 0.9, 300, 100)])
    path = tmp_path / "scan.tiff"
    first = Image.new("RGB", (40, 40), "white")
    first.save(path, save_all=True, append_images=[Image.new("RGB", (40, 40), "white")])
    read = local_ocr.read_text(path)
    assert engine.calls == 2 and read.text.count("Hemoglobin 10.8") == 2


def test_a_sideways_phone_photo_is_turned_upright_before_reading(tmp_path, monkeypatch):
    seen = []
    monkeypatch.setattr(local_ocr, "_ocr_image", lambda image, page: seen.append(image.size) or [])
    path = tmp_path / "photo.jpg"
    image = Image.new("RGB", (80, 40), "white")
    exif = image.getexif()
    exif[0x0112] = 6  # rotated 90 degrees: the camera stored it sideways
    image.save(path, exif=exif)
    local_ocr.read_text(path)
    assert seen == [(40, 80)]


def test_uploads_accept_word_and_tiff_but_not_other_types(tmp_path):
    from fastapi import HTTPException, UploadFile

    from services.file_service import DOCX_TYPE, validate_stored_file, validate_upload

    def upload(name, content_type):
        return UploadFile(file=None, filename=name, headers={"content-type": content_type})

    assert validate_upload(upload("r.docx", DOCX_TYPE)) == ".docx"
    assert validate_upload(upload("scan.tif", "image/tiff")) == ".tif"
    for name, content_type in [
        ("r.doc", "application/msword"),
        ("r.exe", "application/octet-stream"),
        ("r.docx", "text/plain"),
    ]:
        with pytest.raises(HTTPException):
            validate_upload(upload(name, content_type))
    fake = tmp_path / "fake.docx"
    fake.write_bytes(b"not a zip file")
    with pytest.raises(HTTPException):
        validate_stored_file(fake, ".docx")
    assert not fake.exists()  # a rejected file is deleted
    validate_stored_file(make_docx(tmp_path), ".docx")
